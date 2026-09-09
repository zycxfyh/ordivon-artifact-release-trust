#!/usr/bin/env python3
"""Standard-native Artifact Build & Delivery R1 gates.

This module deliberately does not define a universal document model. It checks
build/delivery facts around native artifacts and delegates format conformance to
mature validators whenever available.
"""

from __future__ import annotations

import argparse
import base64
import datetime as dt
import hashlib
import importlib.util
import importlib.metadata
import json
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import subprocess
import sys
import urllib.parse
import xml.etree.ElementTree as ET
import zipfile
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SCHEMA = ROOT / "artifact-delivery/profile-v1.schema.json"
DEFAULT_REQUEST_SCHEMA = ROOT / "artifact-delivery/request-v1.schema.json"
DEFAULT_PRESENTATION_SOURCE_SCHEMA = ROOT / "artifact-delivery/presentation-source-v1.schema.json"
DEFAULT_ATTESTATION_TRUST_POLICY_SCHEMA = ROOT / "artifact-delivery/attestation-trust-policy-v1.schema.json"
DEFAULT_WINDOWS_FONTS = Path("/mnt/c/Windows/Fonts")
IN_TOTO_STATEMENT_V1 = "https://in-toto.io/Statement/v1"
SLSA_PROVENANCE_V1 = "https://slsa.dev/provenance/v1"
SLSA_VERIFICATION_SUMMARY_V1 = "https://slsa.dev/verification_summary/v1"
SLSA_VERSION = "1.2"
LOCAL_VSA_VERIFIER_ID = "https://ordivon.local/verifiers/artifact-delivery-v1"
LOCAL_BUILD_PLATFORM_ID = "https://ordivon.local/build-platforms/workstation-artifact-delivery-v1"
GITHUB_ACTIONS_OIDC_ISSUER = "https://token.actions.githubusercontent.com"
SIGSTORE_BUNDLE_V03 = "application/vnd.dev.sigstore.bundle.v0.3+json"
INTOTO_DSSE_PAYLOAD_TYPE = "application/vnd.in-toto+json"
COSIGN_STANDARD_BUNDLE_MIN_VERSION = (3, 0, 6)
COSIGN_LOCK_PATH = ROOT / "artifact-delivery/toolchain-v1.lock.json"
COSIGN_SELECTED_BINARY = ROOT / ".cache/artifact-toolchain/cosign/current/bin/cosign"
COSIGN_ARCH_PACKAGE = ROOT / ".cache/artifact-toolchain/cosign/bootstrap-arch-3.1.3-1/cosign-3.1.3-1-x86_64.pkg.tar.zst"
COSIGN_ARCH_PACKAGE_SIGNATURE = Path(str(COSIGN_ARCH_PACKAGE) + ".sig")
_COSIGN_PROVENANCE_CACHE: dict[tuple[str, str, str], dict[str, Any]] = {}
VSA_GATE_NAMES = frozenset({
    "profileSchema",
    "structural",
    "dependency",
    "semantic",
    "visual",
    "target",
    "accessibility",
    "conformance",
    "deliveryReadback",
})
ASSEMBLY_GATE_NAMES = frozenset({"companionPdf", "releaseProvenance"})
OOXML_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat().replace("+00:00", "Z")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_fact(path: Path, name: str | None = None) -> dict[str, Any]:
    resolved = path.resolve()
    if not resolved.is_file():
        raise RuntimeError(f"required file is absent: {resolved}")
    return {
        "name": name or resolved.name,
        "path": str(resolved),
        "size": resolved.stat().st_size,
        "digest": {"sha256": sha256_file(resolved)},
    }


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _minimal_profile_checks(value: Any) -> list[str]:
    errors: list[str] = []
    if not isinstance(value, dict):
        return ["profile must be a JSON object"]
    required = {
        "profileVersion",
        "id",
        "artifactClass",
        "authorityMode",
        "locale",
        "primaryOutput",
        "targetRenderer",
        "gates",
        "deliveryTargets",
    }
    missing = sorted(required - set(value))
    if missing:
        errors.append("missing required fields: " + ", ".join(missing))
    if value.get("profileVersion") != 1:
        errors.append("profileVersion must equal 1")
    expected_primary = {
        "presentation": "pptx",
        "document": "docx",
        "spreadsheet": "xlsx",
        "web": "html",
    }
    artifact_class = value.get("artifactClass")
    primary = value.get("primaryOutput")
    if artifact_class in expected_primary:
        if not isinstance(primary, dict) or primary.get("format") != expected_primary[artifact_class]:
            errors.append(f"{artifact_class} primaryOutput.format must be {expected_primary[artifact_class]}")
    if artifact_class in {"fixed-view", "archive", "accessible"}:
        if not isinstance(primary, dict) or primary.get("format") not in {"pdf", "pdf-a-4", "pdf-ua-2"}:
            errors.append(f"{artifact_class} primaryOutput.format must be a PDF format")
    if artifact_class == "presentation":
        for field in ("aspectRatio", "fontPolicy", "fonts"):
            if field not in value:
                errors.append(f"presentation profile requires {field}")
    gates = value.get("gates")
    if not isinstance(gates, dict):
        errors.append("gates must be an object")
    else:
        for key in ("profileSchema", "structural", "target", "deliveryReadback"):
            if not isinstance(gates.get(key), bool):
                errors.append(f"gates.{key} must be boolean")
        for key, flag in gates.items():
            if not isinstance(flag, bool):
                errors.append(f"gates.{key} must be boolean")
    return errors


def validate_json_document(document_path: Path, schema_path: Path, expected_kind: str | None = None) -> dict[str, Any]:
    value = load_json(document_path)
    schema = load_json(schema_path)
    failures: list[str] = []
    if not isinstance(value, dict):
        failures.append("document must be a JSON object")
    elif expected_kind is not None and value.get("kind") != expected_kind:
        failures.append(f"document kind must equal {expected_kind}")
    validator = "jsonschema"
    schema_status = "NOT_RUN"
    schema_error: str | None = None
    if importlib.util.find_spec("jsonschema") is None:
        validator = "unavailable"
        schema_error = "Python jsonschema package is not installed"
    else:
        try:
            import jsonschema  # type: ignore

            jsonschema.Draft202012Validator.check_schema(schema)
            instance = dict(value) if isinstance(value, dict) else value
            if isinstance(instance, dict):
                instance.pop("$schema", None)
            jsonschema.Draft202012Validator(schema, format_checker=jsonschema.FormatChecker()).validate(instance)
            schema_status = "PASS"
        except Exception as error:
            schema_status = "FAIL"
            schema_error = str(error)
            failures.append(f"JSON Schema validation failed: {error}")
    return {
        "status": "PASS" if not failures and schema_status == "PASS" else "FAIL",
        "document": value,
        "failures": failures,
        "jsonSchema": {
            "dialect": "https://json-schema.org/draft/2020-12/schema",
            "validator": validator,
            "status": schema_status,
            "error": schema_error,
            "schemaPath": str(schema_path.resolve()),
        },
    }


def _resolve_request_path(request_path: Path, relative: str) -> Path:
    candidate = Path(relative)
    if candidate.is_absolute():
        return candidate.resolve()
    return (request_path.resolve().parent / candidate).resolve()


def validate_delivery_request(
    request_path: Path,
    request_schema_path: Path = DEFAULT_REQUEST_SCHEMA,
    profile_schema_path: Path = DEFAULT_SCHEMA,
) -> dict[str, Any]:
    request_result = validate_json_document(request_path, request_schema_path, "artifact-delivery-request")
    request = request_result.get("document", {})
    failures = list(request_result.get("failures", []))
    profile_result: dict[str, Any] | None = None
    source_result: dict[str, Any] | None = None
    resolved: dict[str, Any] = {}
    if isinstance(request, dict):
        profile_ref = request.get("profile", {})
        source_ref = request.get("source", {})
        try:
            profile_path = _resolve_request_path(request_path, str(profile_ref.get("path", "")))
            if sha256_file(profile_path) != profile_ref.get("sha256"):
                failures.append("profile digest mismatch")
            profile_result = validate_profile(profile_path, profile_schema_path)
            if profile_result.get("status") != "PASS":
                failures.append("referenced delivery profile did not PASS validation")
            if profile_result.get("profile", {}).get("id") != profile_ref.get("id"):
                failures.append("profile id does not match referenced profile bytes")
            resolved["profile"] = file_fact(profile_path)
        except Exception as error:
            failures.append(f"profile reference error: {error}")
        try:
            source_path = _resolve_request_path(request_path, str(source_ref.get("path", "")))
            if sha256_file(source_path) != source_ref.get("sha256"):
                failures.append("source digest mismatch")
            source_kind = source_ref.get("kind")
            if source_kind == "presentation-source-v1":
                source_result = validate_json_document(source_path, DEFAULT_PRESENTATION_SOURCE_SCHEMA, "presentation-source")
                if source_result.get("status") != "PASS":
                    failures.append("presentation source did not PASS validation")
            resolved["source"] = file_fact(source_path)
        except Exception as error:
            failures.append(f"source reference error: {error}")
        material_results: list[dict[str, Any]] = []
        for item in request.get("materials", []) if isinstance(request.get("materials"), list) else []:
            try:
                material_path = _resolve_request_path(request_path, str(item.get("path", "")))
                actual = sha256_file(material_path)
                if actual != item.get("sha256"):
                    failures.append(f"material digest mismatch: {item.get('path')}")
                material_results.append(file_fact(material_path))
            except Exception as error:
                failures.append(f"material reference error: {error}")
        resolved["materials"] = material_results
    if request_result.get("status") != "PASS":
        failures.append("request envelope did not PASS schema validation")
    return {
        "status": "PASS" if not failures else "FAIL",
        "request": request,
        "requestValidation": request_result,
        "profileValidation": profile_result,
        "sourceValidation": source_result,
        "resolved": resolved,
        "failures": failures,
        "boundary": "Request validation binds exact profile/source/material bytes and schema contracts only. It does not establish build success, target rendering, visual acceptance, accessibility or delivery completion.",
    }


def compile_delivery_plan(request_path: Path) -> dict[str, Any]:
    validation = validate_delivery_request(request_path)
    request = validation.get("request", {})
    profile = (validation.get("profileValidation") or {}).get("profile", {})
    source = request.get("source", {}) if isinstance(request, dict) else {}
    artifact_class = profile.get("artifactClass")
    source_kind = source.get("kind")
    adapter_map = {
        ("presentation", "presentation-source-v1"): {
            "adapter": "python-pptx-presentation-source-v1",
            "buildType": "https://ordivon.local/build-types/artifact-delivery/presentation-source-v1",
        },
        ("document", "markdown"): {
            "adapter": "pandoc-docx",
            "buildType": "https://ordivon.local/build-types/artifact-delivery/markdown-docx-v1",
        },
        ("web", "html-source"): {
            "adapter": "standards-web-source",
            "buildType": "https://ordivon.local/build-types/artifact-delivery/html-source-v1",
        },
    }
    for cls in ("presentation", "document", "spreadsheet", "fixed-view", "archive", "accessible", "web"):
        adapter_map[(cls, "native-file")] = {
            "adapter": "native-artifact-pass-through",
            "buildType": "https://ordivon.local/build-types/artifact-delivery/native-pass-through-v1",
        }
    selection = adapter_map.get((artifact_class, source_kind))
    failures = list(validation.get("failures", []))
    if selection is None:
        failures.append(f"no v1 build adapter for artifactClass={artifact_class!r}, source.kind={source_kind!r}")
    gates = profile.get("gates", {}) if isinstance(profile, dict) else {}
    required_gates = sorted(name for name, required in gates.items() if required is True)
    expected_outputs = [profile.get("primaryOutput")] + list(profile.get("companions", [])) if profile else []
    return {
        "schemaVersion": 1,
        "kind": "artifact-delivery-derived-plan",
        "status": "PASS" if not failures else "FAIL",
        "requestId": request.get("requestId") if isinstance(request, dict) else None,
        "requestSha256": sha256_file(request_path),
        "profileId": profile.get("id") if isinstance(profile, dict) else None,
        "artifactClass": artifact_class,
        "sourceKind": source_kind,
        "buildAdapter": selection["adapter"] if selection else None,
        "buildType": selection["buildType"] if selection else None,
        "resolvedInputs": validation.get("resolved", {}),
        "expectedOutputs": expected_outputs,
        "requiredGates": required_gates,
        "deliveryTargets": list(profile.get("deliveryTargets", [])) if isinstance(profile, dict) else [],
        "stages": ["build", "verify", "package", "release"],
        "failures": failures,
        "boundary": "This is a derived execution projection from immutable request/profile bytes, not durable workflow state. The request does not choose builder.id; execution-platform identity comes from the trusted execution substrate and signer-builder policy. Temporal owns durable retries/timers/admission when execution is wired to the existing workflow substrate.",
    }


def _primary_suffix(profile: dict[str, Any]) -> str:
    suffixes = {
        "pptx": ".pptx",
        "docx": ".docx",
        "xlsx": ".xlsx",
        "html": ".html",
        "pdf": ".pdf",
        "pdf-a-4": ".pdf",
        "pdf-ua-2": ".pdf",
    }
    fmt = str(profile.get("primaryOutput", {}).get("format", ""))
    suffix = suffixes.get(fmt)
    if suffix is None:
        raise RuntimeError(f"unsupported primary output format: {fmt}")
    return suffix


def _request_output_name(request_id: str, suffix: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", request_id).strip("-._") or "artifact"
    return stem + suffix


def execute_build_stage(request_path: Path, output_directory: Path | None = None) -> dict[str, Any]:
    plan = compile_delivery_plan(request_path)
    if plan.get("status") != "PASS":
        return {
            "schemaVersion": 1,
            "kind": "artifact-delivery-build-stage",
            "status": "FAIL",
            "plan": plan,
            "failures": ["derived delivery plan did not PASS"],
        }
    validation = validate_delivery_request(request_path)
    request = validation["request"]
    profile = validation["profileValidation"]["profile"]
    source_path = Path(validation["resolved"]["source"]["path"])
    profile_path = Path(validation["resolved"]["profile"]["path"])
    if output_directory is None:
        output_directory = _resolve_request_path(request_path, str(request["outputDirectory"]))
    output_directory.mkdir(parents=True, exist_ok=True)
    suffix = _primary_suffix(profile)
    output_path = output_directory / _request_output_name(str(request["requestId"]), suffix)
    adapter = plan["buildAdapter"]
    adapter_result: dict[str, Any]
    if adapter == "python-pptx-presentation-source-v1":
        adapter_result = build_presentation_source(source_path, profile_path, output_path)
    elif adapter == "pandoc-docx":
        pandoc = Path(os.environ.get("ARTIFACT_PANDOC", ROOT / ".cache/artifact-toolchain/pandoc/current/bin/pandoc"))
        if not pandoc.is_file():
            adapter_result = {"status": "FAIL", "error": f"Pandoc not found: {pandoc}"}
        else:
            proc = subprocess.run([str(pandoc), str(source_path), "-o", str(output_path)], text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False, timeout=60)
            adapter_result = {
                "status": "PASS" if proc.returncode == 0 and output_path.is_file() else "FAIL",
                "returnCode": proc.returncode,
                "stdout": proc.stdout[-2000:],
                "stderr": proc.stderr[-4000:],
                "artifact": file_fact(output_path) if output_path.is_file() else None,
            }
    elif adapter in {"standards-web-source", "native-artifact-pass-through"}:
        shutil.copyfile(source_path, output_path)
        adapter_result = {
            "status": "PASS" if sha256_file(source_path) == sha256_file(output_path) else "FAIL",
            "source": file_fact(source_path),
            "artifact": file_fact(output_path),
        }
    else:
        adapter_result = {"status": "FAIL", "error": f"unimplemented adapter: {adapter}"}
    failures: list[str] = []
    if adapter_result.get("status") != "PASS":
        failures.append(f"build adapter did not PASS: {adapter}")
    if not output_path.is_file():
        failures.append("primary output file is absent")
    return {
        "schemaVersion": 1,
        "kind": "artifact-delivery-build-stage",
        "status": "PASS" if not failures else "FAIL",
        "request": file_fact(request_path),
        "plan": plan,
        "adapterResult": adapter_result,
        "artifact": file_fact(output_path) if output_path.is_file() else None,
        "failures": failures,
        "boundary": "Build-stage PASS means exact request/profile/source bytes produced the primary artifact through the selected adapter. Independent verification, target rendering, packaging, provenance and release gates are not implied.",
    }


def aggregate_gate_results(profile_path: Path, gate_paths: dict[str, Path]) -> dict[str, Any]:
    profile_result = validate_profile(profile_path)
    profile = profile_result.get("profile", {})
    required = {name for name, flag in profile.get("gates", {}).items() if flag is True}
    components: dict[str, Any] = {}
    failures: list[str] = []
    for gate, path in sorted(gate_paths.items()):
        value = load_json(path)
        status = value.get("status") if isinstance(value, dict) else None
        components[gate] = {
            "status": status,
            "evidence": file_fact(path),
        }
        if gate not in profile.get("gates", {}):
            failures.append(f"gate evidence supplied for undeclared gate: {gate}")
    for gate in sorted(required):
        if gate not in components:
            failures.append(f"required gate evidence missing: {gate}")
        elif components[gate].get("status") != "PASS":
            failures.append(f"required gate did not PASS: {gate}")
    if profile_result.get("status") != "PASS":
        failures.append("profile schema did not PASS")
    return {
        "schemaVersion": 1,
        "kind": "artifact-delivery-gate-aggregation",
        "status": "PASS" if not failures else "FAIL",
        "profileId": profile.get("id"),
        "requiredGates": sorted(required),
        "components": components,
        "failures": failures,
        "boundary": "Aggregation only composes independently-produced gate statuses and evidence digests. It does not execute or reinterpret the underlying validators.",
    }


def _hex_color(value: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(value.upper())


def build_presentation_source(
    source_path: Path,
    profile_path: Path,
    output_path: Path,
    source_schema_path: Path = DEFAULT_PRESENTATION_SOURCE_SCHEMA,
) -> dict[str, Any]:
    source_result = validate_json_document(source_path, source_schema_path, "presentation-source")
    profile_result = validate_profile(profile_path)
    failures: list[str] = []
    source = source_result.get("document", {})
    profile = profile_result.get("profile", {})
    if source_result.get("status") != "PASS":
        failures.append("presentation source schema did not PASS")
    if profile_result.get("status") != "PASS":
        failures.append("delivery profile schema did not PASS")
    if source.get("sourceMode") != "native-composition":
        failures.append("v1 native builder accepts sourceMode=native-composition only")
    if source.get("profileId") != profile.get("id"):
        failures.append("presentation source profileId does not match selected profile")
    if source.get("aspectRatio") != profile.get("aspectRatio"):
        failures.append("presentation source aspectRatio does not match selected profile")
    declared_fonts = {str(item.get("family")) for item in profile.get("fonts", [])}
    width = float(source.get("slideSizeInches", {}).get("width", 0))
    height = float(source.get("slideSizeInches", {}).get("height", 0))
    if width <= 0 or height <= 0:
        failures.append("presentation slide size must be positive")
    slide_ids: set[str] = set()
    for slide in source.get("slides", []) if isinstance(source.get("slides"), list) else []:
        slide_id = str(slide.get("id"))
        if slide_id in slide_ids:
            failures.append(f"duplicate slide id: {slide_id}")
        slide_ids.add(slide_id)
        element_ids: set[str] = set()
        for element in slide.get("elements", []) if isinstance(slide.get("elements"), list) else []:
            element_id = str(element.get("id"))
            if element_id in element_ids:
                failures.append(f"duplicate element id on {slide_id}: {element_id}")
            element_ids.add(element_id)
            family = str(element.get("fontFamily"))
            if family not in declared_fonts:
                failures.append(f"undeclared font family on {slide_id}/{element_id}: {family}")
            box = element.get("box", {})
            x, y, w, h = (float(box.get(key, 0)) for key in ("x", "y", "w", "h"))
            if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > width + 1e-6 or y + h > height + 1e-6:
                failures.append(f"out-of-bounds box on {slide_id}/{element_id}")
    if failures:
        return {
            "status": "FAIL",
            "source": file_fact(source_path),
            "profile": file_fact(profile_path),
            "failures": failures,
        }
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    align_map = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }
    prs = Presentation()
    prs.slide_width = Inches(width)
    prs.slide_height = Inches(height)
    blank = prs.slide_layouts[6]
    for slide_spec in source.get("slides", []):
        slide = prs.slides.add_slide(blank)
        for element in slide_spec.get("elements", []):
            if element.get("kind") != "text":
                raise RuntimeError(f"unsupported presentation element kind: {element.get('kind')}")
            box = element["box"]
            shape = slide.shapes.add_textbox(Inches(box["x"]), Inches(box["y"]), Inches(box["w"]), Inches(box["h"]))
            frame = shape.text_frame
            frame.clear()
            paragraph = frame.paragraphs[0]
            paragraph.text = element.get("text", "")
            paragraph.alignment = align_map.get(element.get("align", "left"), PP_ALIGN.LEFT)
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            run.font.name = element["fontFamily"]
            run.font.size = Pt(float(element["fontSizePt"]))
            run.font.bold = bool(element.get("bold", False))
            run.font.italic = bool(element.get("italic", False))
            if element.get("colorHex"):
                run.font.color.rgb = _hex_color(str(element["colorHex"]))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    built = inspect_pptx(output_path, profile.get("semanticPolicy", {}).get("placeholderPatterns", []))
    semantic = verify_presentation_semantics(profile, built)
    post_failures: list[str] = []
    if built.get("status") != "PASS":
        post_failures.append("built PPTX failed package/relationship inspection")
    if semantic.get("status") != "PASS":
        post_failures.append("built PPTX failed presentation semantic checks")
    return {
        "status": "PASS" if not post_failures else "FAIL",
        "source": file_fact(source_path),
        "profile": file_fact(profile_path),
        "artifact": file_fact(output_path),
        "presentationId": source.get("presentationId"),
        "slideCount": len(source.get("slides", [])),
        "builder": {"implementation": "python-pptx", "version": importlib.metadata.version("python-pptx")},
        "inspection": built,
        "semantic": semantic,
        "failures": post_failures,
        "boundary": "Builder PASS establishes source/profile binding plus native PPTX package/semantic checks. Open XML SDK, target PowerPoint, visual, accessibility and delivery gates remain independent.",
    }


def validate_profile(profile_path: Path, schema_path: Path = DEFAULT_SCHEMA) -> dict[str, Any]:
    profile = load_json(profile_path)
    schema = load_json(schema_path)
    errors = _minimal_profile_checks(profile)
    validator = "jsonschema"
    schema_status = "NOT_RUN"
    schema_error: str | None = None
    if importlib.util.find_spec("jsonschema") is None:
        validator = "unavailable"
        schema_error = "Python jsonschema package is not installed"
    else:
        try:
            import jsonschema  # type: ignore

            jsonschema.Draft202012Validator.check_schema(schema)
            instance = dict(profile)
            instance.pop("$schema", None)
            jsonschema.Draft202012Validator(schema).validate(instance)
            schema_status = "PASS"
        except Exception as error:  # pragma: no cover - depends on optional validator
            schema_status = "FAIL"
            schema_error = str(error)
            errors.append(f"JSON Schema validation failed: {error}")
    return {
        "status": "PASS" if not errors and schema_status == "PASS" else "FAIL",
        "profile": profile,
        "minimalContractErrors": errors,
        "jsonSchema": {
            "dialect": "https://json-schema.org/draft/2020-12/schema",
            "validator": validator,
            "status": schema_status,
            "error": schema_error,
            "schemaPath": str(schema_path.resolve()),
        },
    }


def _safe_zip_names(names: Iterable[str]) -> list[str]:
    bad: list[str] = []
    for name in names:
        p = PurePosixPath(name)
        if p.is_absolute() or ".." in p.parts:
            bad.append(name)
    return bad


def _relationship_base(rels_name: str) -> str:
    if rels_name == "_rels/.rels":
        return ""
    marker = "/_rels/"
    if marker not in rels_name or not rels_name.endswith(".rels"):
        return ""
    left, right = rels_name.split(marker, 1)
    owner = posixpath.join(left, right[:-5])
    return posixpath.dirname(owner)


def _resolve_relationship_target(base: str, target: str) -> str | None:
    target = target.split("#", 1)[0]
    parsed = urllib.parse.urlparse(target)
    if parsed.scheme:
        return None
    if target.startswith("/"):
        resolved = posixpath.normpath(target.lstrip("/"))
    else:
        resolved = posixpath.normpath(posixpath.join(base, target))
    if resolved.startswith("../") or resolved == "..":
        return None
    return resolved


def inspect_pptx(path: Path, placeholder_patterns: Iterable[str] = ()) -> dict[str, Any]:
    artifact = file_fact(path)
    failures: list[str] = []
    warnings: list[str] = []
    unresolved_relationships: list[dict[str, str]] = []
    font_names: set[str] = set()
    render_explicit_font_names: set[str] = set()
    placeholder_hits: list[dict[str, str]] = []
    slide_count = 0
    hidden_slides = 0
    slide_size: dict[str, int] | None = None

    if not zipfile.is_zipfile(path):
        return {"status": "FAIL", "artifact": artifact, "failures": ["not a ZIP/OPC package"]}

    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        unsafe = _safe_zip_names(names)
        if unsafe:
            failures.append(f"unsafe package paths: {unsafe[:5]}")
        required = {
            "[Content_Types].xml",
            "_rels/.rels",
            "ppt/presentation.xml",
            "ppt/_rels/presentation.xml.rels",
        }
        missing = sorted(required - names)
        if missing:
            failures.append("missing required OPC/PPTX parts: " + ", ".join(missing))

        for rels_name in sorted(n for n in names if n.endswith(".rels")):
            try:
                root = ET.fromstring(package.read(rels_name))
            except Exception as error:
                failures.append(f"invalid relationships XML {rels_name}: {error}")
                continue
            base = _relationship_base(rels_name)
            for rel in root.findall(f"{{{OOXML_REL_NS}}}Relationship"):
                if rel.attrib.get("TargetMode") == "External":
                    continue
                target = rel.attrib.get("Target", "")
                resolved = _resolve_relationship_target(base, target)
                if resolved is None:
                    warnings.append(f"unresolved non-file relationship target in {rels_name}: {target}")
                    continue
                if resolved not in names:
                    unresolved_relationships.append({"rels": rels_name, "target": target, "resolved": resolved})
        if unresolved_relationships:
            failures.append(f"{len(unresolved_relationships)} internal relationship target(s) are missing")

        if "ppt/presentation.xml" in names:
            try:
                root = ET.fromstring(package.read("ppt/presentation.xml"))
                ids = root.findall(f".//{{{PRESENTATION_NS}}}sldId")
                slide_count = len(ids)
                size_node = root.find(f".//{{{PRESENTATION_NS}}}sldSz")
                if size_node is not None:
                    try:
                        cx = int(size_node.attrib.get("cx", "0"))
                        cy = int(size_node.attrib.get("cy", "0"))
                        if cx > 0 and cy > 0:
                            slide_size = {"cx": cx, "cy": cy}
                    except (TypeError, ValueError):
                        pass
            except Exception as error:
                failures.append(f"invalid ppt/presentation.xml: {error}")

        patterns = [re.compile(p, re.IGNORECASE) for p in placeholder_patterns]
        render_xml_names = sorted(
            n
            for n in names
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
            or re.fullmatch(r"ppt/slideLayouts/slideLayout\d+\.xml", n)
            or re.fullmatch(r"ppt/slideMasters/slideMaster\d+\.xml", n)
        )
        for xml_name in render_xml_names:
            try:
                root = ET.fromstring(package.read(xml_name))
            except Exception as error:
                failures.append(f"invalid render-relevant XML {xml_name}: {error}")
                continue
            is_slide = bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", xml_name))
            if is_slide and root.attrib.get("show") in {"0", "false", "False"}:
                hidden_slides += 1
            if is_slide:
                texts = [node.text or "" for node in root.findall(f".//{{{DRAWING_NS}}}t")]
                joined = "\n".join(texts)
                for pattern in patterns:
                    match = pattern.search(joined)
                    if match:
                        placeholder_hits.append({"slide": xml_name, "pattern": pattern.pattern, "match": match.group(0)})
            for node in root.iter():
                typeface = node.attrib.get("typeface")
                if typeface and not typeface.startswith("+"):
                    font_names.add(typeface)
                    render_explicit_font_names.add(typeface)

        for xml_name in sorted(n for n in names if n.startswith("ppt/theme/") and n.endswith(".xml")):
            try:
                root = ET.fromstring(package.read(xml_name))
            except Exception:
                continue
            for node in root.iter():
                typeface = node.attrib.get("typeface")
                if typeface and not typeface.startswith("+"):
                    font_names.add(typeface)

    if slide_count <= 0:
        failures.append("presentation has no slides")
    if placeholder_hits:
        failures.append(f"placeholder text found in {len(placeholder_hits)} slide occurrence(s)")
    return {
        "status": "PASS" if not failures else "FAIL",
        "artifact": artifact,
        "package": {
            "kind": "OOXML/OPC PPTX",
            "slideCount": slide_count,
            "hiddenSlideCount": hidden_slides,
            "slideSizeEmu": slide_size,
            "aspectRatio": (round(slide_size["cx"] / slide_size["cy"], 8) if slide_size else None),
            "referencedTypefaceNames": sorted(font_names),
            "renderExplicitTypefaceNames": sorted(render_explicit_font_names),
            "unresolvedRelationships": unresolved_relationships,
            "placeholderHits": placeholder_hits,
        },
        "scope": {
            "packageChecks": "performed",
            "OOXMLSchemaValidation": "NOT_RUN",
            "note": "Package/relationship/XML checks do not replace an ISO/IEC 29500 schema validator such as Open XML SDK validation.",
        },
        "warnings": warnings,
        "failures": failures,
    }


def verify_openxml_evidence(evidence_path: Path, artifact: Path) -> dict[str, Any]:
    value = load_json(evidence_path)
    failures: list[str] = []
    expected_digest = sha256_file(artifact)
    if value.get("artifact", {}).get("sha256") != expected_digest:
        failures.append("Open XML validation evidence artifact digest mismatch")
    validator = value.get("validator", {})
    if validator.get("implementation") != "DocumentFormat.OpenXml":
        failures.append("Open XML validation evidence did not use DocumentFormat.OpenXml")
    if validator.get("api") != "OpenXmlValidator":
        failures.append("Open XML validation evidence did not use OpenXmlValidator")
    if not validator.get("packageVersion"):
        failures.append("Open XML validation evidence omitted package version")
    if value.get("status") != "PASS":
        failures.append("Open XML validator did not PASS")
    if int(value.get("validationErrorCount", -1)) != 0:
        failures.append("Open XML validator reported validation errors")
    return {
        "status": "PASS" if not failures else "FAIL",
        "evidencePath": str(evidence_path.resolve()),
        "validator": validator,
        "failures": failures,
    }


def verify_presentation_semantics(profile: dict[str, Any], pptx_result: dict[str, Any]) -> dict[str, Any]:
    failures: list[str] = []
    package = pptx_result.get("package", {})
    slide_count = int(package.get("slideCount", 0))
    policy = profile.get("semanticPolicy", {})
    minimum = policy.get("minimumSlideCount")
    maximum = policy.get("maximumSlideCount")
    if isinstance(minimum, int) and slide_count < minimum:
        failures.append(f"slide count {slide_count} is below minimum {minimum}")
    if isinstance(maximum, int) and slide_count > maximum:
        failures.append(f"slide count {slide_count} exceeds maximum {maximum}")
    declared_aspect = profile.get("aspectRatio")
    observed = package.get("aspectRatio")
    target_ratios = {"16:9": 16 / 9, "4:3": 4 / 3}
    if declared_aspect in target_ratios:
        if not isinstance(observed, (int, float)):
            failures.append("presentation slide size/aspect ratio is unavailable")
        elif abs(float(observed) - target_ratios[declared_aspect]) > 0.002:
            failures.append(f"presentation aspect ratio {observed:.6f} does not satisfy profile {declared_aspect}")
    return {
        "status": "PASS" if not failures else "FAIL",
        "slideCount": slide_count,
        "minimumSlideCount": minimum,
        "maximumSlideCount": maximum,
        "declaredAspectRatio": declared_aspect,
        "observedAspectRatio": observed,
        "slideSizeEmu": package.get("slideSizeEmu"),
        "failures": failures,
    }


def verify_font_manifest(
    profile: dict[str, Any],
    font_dir: Path,
    observed_typefaces: Iterable[str] = (),
) -> dict[str, Any]:
    results: list[dict[str, Any]] = []
    failures: list[str] = []
    declared = {str(font.get("family", "")).casefold(): str(font.get("family", "")) for font in profile.get("fonts", [])}
    observed = sorted({str(name) for name in observed_typefaces if str(name)})
    undeclared = [name for name in observed if name.casefold() not in declared]
    for name in undeclared:
        failures.append(f"artifact render graph directly references undeclared typeface: {name}")
    for font in profile.get("fonts", []):
        family = str(font.get("family", ""))
        required = bool(font.get("required"))
        files: list[dict[str, Any]] = []
        for filename in font.get("targetFiles", []):
            path = font_dir / str(filename)
            present = path.is_file() and not path.is_symlink()
            fact = {
                "name": str(filename),
                "path": str(path),
                "present": present,
                "sha256": sha256_file(path) if present else None,
            }
            files.append(fact)
            if required and not present:
                failures.append(f"required font file missing for {family}: {filename}")
        if required and not font.get("targetFiles"):
            failures.append(f"required font {family} has no targetFiles binding")
        results.append(
            {
                "family": family,
                "required": required,
                "embeddingPermission": font.get("embeddingPermission"),
                "files": files,
            }
        )
    return {
        "status": "PASS" if not failures else "FAIL",
        "fontDirectory": str(font_dir),
        "observedRenderExplicitTypefaces": observed,
        "undeclaredObservedTypefaces": undeclared,
        "fonts": results,
        "failures": failures,
    }


def verify_pdf(path: Path) -> dict[str, Any]:
    artifact = file_fact(path)
    qpdf = shutil.which("qpdf")
    if not qpdf:
        return {
            "status": "FAIL",
            "artifact": artifact,
            "validator": "qpdf",
            "error": "qpdf is not installed",
            "profileValidation": "NOT_RUN",
        }
    proc = subprocess.run(
        [qpdf, "--check", str(path)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=30,
    )
    return {
        "status": "PASS" if proc.returncode == 0 else "FAIL",
        "artifact": artifact,
        "validator": qpdf,
        "exitCode": proc.returncode,
        "stdout": proc.stdout.strip()[:4000],
        "stderr": proc.stderr.strip()[:4000],
        "profileValidation": "NOT_RUN",
        "note": "qpdf structural checking does not establish PDF/A or PDF/UA conformance; use veraPDF/PAC in those profiles.",
    }


def _verapdf_executable() -> Path | None:
    configured = os.environ.get("ARTIFACT_VERAPDF")
    if configured:
        path = Path(configured)
        return path if path.is_file() else None
    system = shutil.which("verapdf")
    if system:
        return Path(system)
    local = ROOT / ".cache/artifact-toolchain/verapdf/current/verapdf"
    return local if local.is_file() else None


def verify_pdf_conformance(path: Path, flavour: str) -> dict[str, Any]:
    artifact = file_fact(path)
    if flavour not in {"4", "4f", "4e", "ua1", "ua2", "wt1r", "wt1a"}:
        return {"status": "FAIL", "artifact": artifact, "flavour": flavour, "error": "unsupported veraPDF flavour"}
    executable = _verapdf_executable()
    if executable is None:
        return {"status": "NOT_RUN", "artifact": artifact, "flavour": flavour, "error": "veraPDF is not installed"}
    proc = subprocess.run(
        [str(executable), "--format", "json", "--flavour", flavour, str(path)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=60,
    )
    compliant = False
    profile_name: str | None = None
    failed_rules: int | None = None
    failed_checks: int | None = None
    parse_error: str | None = None
    try:
        parsed = json.loads(proc.stdout)
        validation = parsed["report"]["jobs"][0]["validationResult"][0]
        compliant = validation.get("compliant") is True
        profile_name = validation.get("profileName")
        details = validation.get("details", {})
        failed_rules = details.get("failedRules")
        failed_checks = details.get("failedChecks")
    except Exception as error:
        parse_error = str(error)
    return {
        "status": "PASS" if compliant and proc.returncode == 0 else "FAIL",
        "artifact": artifact,
        "validator": {"implementation": "veraPDF", "executable": str(executable), "flavour": flavour},
        "profileName": profile_name,
        "compliant": compliant,
        "failedRules": failed_rules,
        "failedChecks": failed_checks,
        "exitCode": proc.returncode,
        "parseError": parse_error,
        "stderr": proc.stderr.strip()[:4000],
        "boundary": "veraPDF conformance is machine-checkable profile evidence only; human accessibility/use review and target-viewer acceptance remain separate gates.",
    }


def verify_openxml_artifact(path: Path) -> dict[str, Any]:
    artifact = file_fact(path)
    dotnet = Path(os.environ.get("ARTIFACT_DOTNET", ROOT / ".cache/dotnet/dotnet"))
    validator_dll = ROOT / "artifact-delivery/openxml-validator/bin/Release/net8.0/ArtifactOpenXmlValidator.dll"
    if not dotnet.is_file() or not validator_dll.is_file():
        return {
            "status": "NOT_RUN",
            "artifact": artifact,
            "error": "DocumentFormat.OpenXml validator runtime is unavailable",
        }
    proc = subprocess.run(
        [str(dotnet), str(validator_dll), str(path)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=60,
    )
    parsed: dict[str, Any] | None = None
    parse_error: str | None = None
    try:
        parsed = json.loads(proc.stdout)
    except Exception as error:
        parse_error = str(error)
    return {
        "status": "PASS" if proc.returncode == 0 and parsed and parsed.get("status") == "PASS" else "FAIL",
        "artifact": artifact,
        "validatorOutput": parsed,
        "exitCode": proc.returncode,
        "parseError": parse_error,
        "stderr": proc.stderr.strip()[:4000],
        "boundary": "DocumentFormat.OpenXml schema/semantic validation only; Office target rendering, visual acceptance and delivery remain independent.",
    }


def _vnu_jar() -> Path | None:
    configured = os.environ.get("ARTIFACT_VNU")
    if configured:
        path = Path(configured)
        return path if path.is_file() else None
    local = ROOT / ".cache/artifact-toolchain/vnu/vnu.jar"
    return local if local.is_file() else None


def verify_html_conformance(path: Path) -> dict[str, Any]:
    artifact = file_fact(path)
    jar = _vnu_jar()
    java = shutil.which("java")
    if jar is None or java is None:
        return {
            "status": "NOT_RUN",
            "artifact": artifact,
            "error": "Nu Html Checker or Java is unavailable",
        }
    proc = subprocess.run(
        [java, "-jar", str(jar), "--format", "json", str(path)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=60,
    )
    parsed: dict[str, Any] | None = None
    messages: list[Any] = []
    version: str | None = None
    parse_error: str | None = None
    try:
        payload = proc.stdout.strip() or proc.stderr.strip()
        parsed = json.loads(payload)
        messages = parsed.get("messages", []) if isinstance(parsed, dict) else []
        version = parsed.get("version") if isinstance(parsed, dict) else None
    except Exception as error:
        parse_error = str(error)
    return {
        "status": "PASS" if proc.returncode == 0 and parsed is not None and not messages else "FAIL",
        "artifact": artifact,
        "validator": {
            "implementation": "Nu Html Checker",
            "version": version,
            "jar": str(jar),
            "jarSha256": sha256_file(jar),
        },
        "messageCount": len(messages),
        "messages": messages[:100],
        "exitCode": proc.returncode,
        "parseError": parse_error,
        "stderr": proc.stderr.strip()[:4000],
        "boundary": "Nu Html Checker conformance evidence covers HTML/CSS/SVG syntax/content-model checks; browser behavior, accessibility and deployed-origin behavior remain independent.",
    }


def verify_web_local(path: Path) -> dict[str, Any]:
    artifact = file_fact(path)
    node = shutil.which("node")
    verifier = ROOT / "artifact-delivery/node/verify_html.mjs"
    if node is None or not verifier.is_file():
        return {"status": "NOT_RUN", "artifact": artifact, "error": "Node/Playwright HTML verifier is unavailable"}
    proc = subprocess.run(
        [node, str(verifier), str(path.resolve())],
        cwd=ROOT / "artifact-delivery/node",
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=90,
    )
    parsed: dict[str, Any] | None = None
    parse_error: str | None = None
    try:
        parsed = json.loads(proc.stdout)
    except Exception as error:
        parse_error = str(error)
    digest_ok = parsed is not None and parsed.get("subject", {}).get("sha256") == artifact["digest"]["sha256"]
    return {
        "status": "PASS" if proc.returncode == 0 and parsed and parsed.get("status") == "PASS" and digest_ok else "FAIL",
        "artifact": artifact,
        "verifierOutput": parsed,
        "digestBound": digest_ok,
        "exitCode": proc.returncode,
        "parseError": parse_error,
        "stderr": proc.stderr.strip()[:4000],
        "boundary": "Local Playwright/axe evidence only; delivery profile policy decides which renderers are required and unsupported-host WebKit cannot be promoted to PASS.",
    }


def ni_sha256_uri(path: Path) -> str:
    digest = hashlib.sha256(path.read_bytes()).digest()
    value = base64.urlsafe_b64encode(digest).decode("ascii").rstrip("=")
    return f"ni:///sha-256;{value}"


def verification_summary_statement(
    subject: Path,
    profile_path: Path,
    verifier_id: str,
    verifier_versions: dict[str, str],
    passed: bool,
) -> dict[str, Any]:
    verifier_id = _require_uri(verifier_id, "verifier-id")
    profile_uri = profile_path.resolve().as_uri()
    subject_uri = ni_sha256_uri(subject)
    return {
        "_type": IN_TOTO_STATEMENT_V1,
        "subject": [{"name": subject.name, "digest": {"sha256": sha256_file(subject)}}],
        "predicateType": SLSA_VERIFICATION_SUMMARY_V1,
        "predicate": {
            "verifier": {"id": verifier_id, "version": dict(sorted(verifier_versions.items()))},
            "timeVerified": utc_now(),
            "resourceUri": subject_uri,
            "policy": {
                "uri": profile_uri,
                "digest": {"sha256": sha256_file(profile_path)},
            },
            "verificationResult": "PASSED" if passed else "FAILED",
            "verifiedLevels": ["SLSA_BUILD_LEVEL_UNEVALUATED" if passed else "FAILED"],
            "slsaVersion": SLSA_VERSION,
        },
    }


def verify_verification_summary(
    statement_path: Path,
    subject: Path,
    profile_path: Path,
    allowed_verifiers: Iterable[str] = (),
) -> dict[str, Any]:
    value = load_json(statement_path)
    failures: list[str] = []
    expected_subject = {subject.name: sha256_file(subject)}
    observed_subject: dict[str, str] = {}
    if value.get("_type") != IN_TOTO_STATEMENT_V1:
        failures.append("verification summary is not an in-toto Statement v1")
    if value.get("predicateType") != SLSA_VERIFICATION_SUMMARY_V1:
        failures.append("predicateType is not SLSA Verification Summary v1")
    for item in value.get("subject", []) if isinstance(value.get("subject"), list) else []:
        if isinstance(item, dict) and isinstance(item.get("name"), str):
            digest = item.get("digest", {}).get("sha256") if isinstance(item.get("digest"), dict) else None
            if isinstance(digest, str):
                observed_subject[item["name"]] = digest
    if observed_subject != expected_subject:
        failures.append("verification summary subject does not exactly bind the artifact")
    predicate = value.get("predicate", {}) if isinstance(value.get("predicate"), dict) else {}
    verifier = predicate.get("verifier", {}) if isinstance(predicate.get("verifier"), dict) else {}
    verifier_id = verifier.get("id")
    if not isinstance(verifier_id, str) or not urllib.parse.urlparse(verifier_id).scheme:
        failures.append("verification summary verifier.id is absent or not a URI")
    allowed = set(allowed_verifiers)
    if allowed and verifier_id not in allowed:
        failures.append("verification summary verifier.id is not allowed for this gate")
    policy = predicate.get("policy", {}) if isinstance(predicate.get("policy"), dict) else {}
    if policy.get("uri") != profile_path.resolve().as_uri():
        failures.append("verification summary policy URI does not bind the selected delivery profile")
    if policy.get("digest", {}).get("sha256") != sha256_file(profile_path):
        failures.append("verification summary policy digest does not bind the selected delivery profile")
    if predicate.get("resourceUri") != ni_sha256_uri(subject):
        failures.append("verification summary resourceUri does not bind the artifact content via RFC 6920 ni URI")
    input_attestations = predicate.get("inputAttestations")
    if input_attestations not in (None, []):
        failures.append("local v1 verification summaries must not misuse inputAttestations for raw validator output")

    verification_result = predicate.get("verificationResult")
    if verification_result not in {"PASSED", "FAILED"}:
        failures.append("verificationResult must be PASSED or FAILED")
    levels = predicate.get("verifiedLevels")
    if not isinstance(levels, list) or not levels:
        failures.append("verifiedLevels is absent or empty")
    elif verification_result == "PASSED" and "SLSA_BUILD_LEVEL_UNEVALUATED" not in levels:
        failures.append("PASSED non-SLSA policy verification must remain BUILD_LEVEL_UNEVALUATED")
    if predicate.get("slsaVersion") != SLSA_VERSION:
        failures.append(f"unexpected slsaVersion; expected {SLSA_VERSION}")
    return {
        "status": "PASS" if not failures else "FAIL",
        "statement": file_fact(statement_path),
        "subject": file_fact(subject),
        "profile": file_fact(profile_path),
        "verifier": verifier,
        "verificationResult": verification_result,
        "inputAttestations": input_attestations,
        "authenticity": "NOT_VERIFIED",
        "failures": failures,
        "boundary": "This validates unsigned in-toto/SLSA VSA structure and exact local digest bindings. External trust still requires signature/root-of-trust verification; do not treat this result as cryptographic authenticity.",
    }


def _resolve_policy_path(policy_path: Path, value: str) -> Path:
    candidate = Path(value)
    return candidate.resolve() if candidate.is_absolute() else (policy_path.resolve().parent / candidate).resolve()


def _minimal_attestation_trust_policy_checks(value: Any) -> list[str]:
    failures: list[str] = []
    if not isinstance(value, dict):
        return ["attestation trust policy must be a JSON object"]
    if value.get("policyVersion") != 1:
        failures.append("policyVersion must equal 1")
    accepted = value.get("acceptedBundleMediaTypes")
    if accepted != [SIGSTORE_BUNDLE_V03] and not (
        isinstance(accepted, list) and accepted and set(accepted) == {SIGSTORE_BUNDLE_V03}
    ):
        failures.append("acceptedBundleMediaTypes must contain only the standardized Sigstore v0.3 JSON bundle media type")
    signers = value.get("signers")
    if not isinstance(signers, list) or not signers:
        failures.append("signers must be a non-empty array")
        return failures
    seen: set[str] = set()
    for signer in signers:
        if not isinstance(signer, dict):
            failures.append("every signer must be an object")
            continue
        signer_id = signer.get("id")
        if not isinstance(signer_id, str) or not signer_id:
            failures.append("every signer requires a non-empty id")
        elif signer_id in seen:
            failures.append(f"duplicate signer id: {signer_id}")
        else:
            seen.add(signer_id)
        mode = signer.get("mode")
        if mode not in {"public-key", "keyless"}:
            failures.append(f"unsupported signer mode for {signer_id}: {mode}")
        allowed_verifiers = signer.get("allowedVerifierIds")
        allowed_builders = signer.get("allowedBuilderIds")
        if not isinstance(allowed_verifiers, list) and not isinstance(allowed_builders, list):
            failures.append(f"signer {signer_id} must authorize at least one verifier.id or builder.id")
        for field, allowed in (("allowedVerifierIds", allowed_verifiers), ("allowedBuilderIds", allowed_builders)):
            if allowed is None:
                continue
            if not isinstance(allowed, list) or not allowed:
                failures.append(f"signer {signer_id} requires non-empty {field} when present")
                continue
            for authority_id in allowed:
                if not isinstance(authority_id, str) or not urllib.parse.urlparse(authority_id).scheme:
                    failures.append(f"signer {signer_id} contains a {field} value that is not a URI")
        if not isinstance(signer.get("requireTransparencyLog"), bool):
            failures.append(f"signer {signer_id} requires boolean requireTransparencyLog")
        if mode == "public-key":
            public_key = signer.get("publicKey")
            if not isinstance(public_key, dict) or not isinstance(public_key.get("path"), str):
                failures.append(f"public-key signer {signer_id} requires publicKey.path")
            if not isinstance(public_key, dict) or not re.fullmatch(r"[0-9a-f]{64}", str(public_key.get("sha256", ""))):
                failures.append(f"public-key signer {signer_id} requires lowercase publicKey.sha256")
        if mode == "keyless":
            if signer.get("requireTransparencyLog") is not True:
                failures.append(f"keyless signer {signer_id} must require transparency-log verification")
            if not isinstance(signer.get("certificateIdentity"), str) or not signer.get("certificateIdentity"):
                failures.append(f"keyless signer {signer_id} requires certificateIdentity")
            issuer = signer.get("certificateOidcIssuer")
            if not isinstance(issuer, str) or not urllib.parse.urlparse(issuer).scheme:
                failures.append(f"keyless signer {signer_id} requires a URI certificateOidcIssuer")
            github = signer.get("githubActions")
            if issuer == GITHUB_ACTIONS_OIDC_ISSUER:
                if not isinstance(github, dict):
                    failures.append(f"GitHub Actions keyless signer {signer_id} requires exact githubActions workflow claims")
                else:
                    repository = github.get("repository")
                    workflow_ref = github.get("workflowRef")
                    ref = github.get("ref")
                    sha_authority = github.get("shaAuthority")
                    sha = github.get("sha")
                    name = github.get("name")
                    trigger = github.get("trigger")
                    if not isinstance(repository, str) or not re.fullmatch(r"[^/\s]+/[^/\s]+", repository):
                        failures.append(f"GitHub Actions keyless signer {signer_id} requires owner/repository")
                    if not isinstance(workflow_ref, str) or not re.fullmatch(r"[^/\s]+/[^/\s]+/\.github/workflows/[^@\s]+@.+", workflow_ref):
                        failures.append(f"GitHub Actions keyless signer {signer_id} requires exact workflowRef")
                    if not isinstance(ref, str) or not ref:
                        failures.append(f"GitHub Actions keyless signer {signer_id} requires exact ref")
                    if sha_authority not in {"policy", "runtime"}:
                        failures.append(f"GitHub Actions keyless signer {signer_id} requires shaAuthority=policy|runtime")
                    elif sha_authority == "policy":
                        if not isinstance(sha, str) or not re.fullmatch(r"[0-9a-f]{40}", sha):
                            failures.append(f"GitHub Actions keyless signer {signer_id} with policy SHA authority requires exact 40-hex workflow sha")
                    elif sha is not None:
                        failures.append(f"GitHub Actions keyless signer {signer_id} with runtime SHA authority must not embed githubActions.sha")
                    if not isinstance(name, str) or not name:
                        failures.append(f"GitHub Actions keyless signer {signer_id} requires exact workflow name")
                    if not isinstance(trigger, str) or not trigger:
                        failures.append(f"GitHub Actions keyless signer {signer_id} requires exact workflow trigger")
                    identity = signer.get("certificateIdentity")
                    if isinstance(workflow_ref, str) and identity != f"https://github.com/{workflow_ref}":
                        failures.append(f"GitHub Actions keyless signer {signer_id} certificateIdentity must exactly match githubActions.workflowRef")
                    if isinstance(repository, str) and isinstance(workflow_ref, str) and not workflow_ref.startswith(repository + "/.github/workflows/"):
                        failures.append(f"GitHub Actions keyless signer {signer_id} workflowRef repository does not match githubActions.repository")
                    if isinstance(ref, str) and isinstance(workflow_ref, str) and not workflow_ref.endswith("@" + ref):
                        failures.append(f"GitHub Actions keyless signer {signer_id} workflowRef ref does not match githubActions.ref")
            trusted_root = signer.get("trustedRoot")
            if trusted_root is not None:
                if not isinstance(trusted_root, dict) or not isinstance(trusted_root.get("path"), str):
                    failures.append(f"keyless signer {signer_id} trustedRoot requires path")
                if not isinstance(trusted_root, dict) or not re.fullmatch(r"[0-9a-f]{64}", str(trusted_root.get("sha256", ""))):
                    failures.append(f"keyless signer {signer_id} trustedRoot requires lowercase sha256")
    return failures


def validate_attestation_trust_policy(
    policy_path: Path,
    schema_path: Path = DEFAULT_ATTESTATION_TRUST_POLICY_SCHEMA,
) -> dict[str, Any]:
    result = validate_json_document(policy_path, schema_path)
    policy = result.get("document", {})
    failures = list(result.get("failures", []))
    failures.extend(_minimal_attestation_trust_policy_checks(policy))
    resolved_signers: dict[str, Any] = {}
    if isinstance(policy, dict):
        for signer in policy.get("signers", []) if isinstance(policy.get("signers"), list) else []:
            if not isinstance(signer, dict) or not isinstance(signer.get("id"), str):
                continue
            resolved = dict(signer)
            if signer.get("mode") == "public-key" and isinstance(signer.get("publicKey"), dict):
                public_key = signer["publicKey"]
                if isinstance(public_key.get("path"), str):
                    key_path = _resolve_policy_path(policy_path, public_key["path"])
                    key_fact: dict[str, Any] | None = None
                    if not key_path.is_file():
                        failures.append(f"trusted public key is absent for signer {signer['id']}: {key_path}")
                    else:
                        key_fact = file_fact(key_path)
                        if key_fact["digest"]["sha256"] != public_key.get("sha256"):
                            failures.append(f"trusted public key digest mismatch for signer {signer['id']}")
                    resolved["publicKeyResolved"] = key_fact
                    resolved["publicKeyPath"] = str(key_path)
            trusted_root = signer.get("trustedRoot")
            if isinstance(trusted_root, dict) and isinstance(trusted_root.get("path"), str):
                root_path = _resolve_policy_path(policy_path, trusted_root["path"])
                root_fact: dict[str, Any] | None = None
                if not root_path.is_file():
                    failures.append(f"trustedRoot is absent for signer {signer['id']}: {root_path}")
                else:
                    root_fact = file_fact(root_path)
                    if root_fact["digest"]["sha256"] != trusted_root.get("sha256"):
                        failures.append(f"trustedRoot digest mismatch for signer {signer['id']}")
                resolved["trustedRootResolved"] = root_fact
                resolved["trustedRootPath"] = str(root_path)
            resolved_signers[signer["id"]] = resolved
    schema_ok = result.get("jsonSchema", {}).get("status") == "PASS"
    return {
        "status": "PASS" if schema_ok and not failures else "FAIL",
        "policy": policy,
        "policyFact": file_fact(policy_path),
        "resolvedSigners": resolved_signers,
        "jsonSchema": result.get("jsonSchema"),
        "failures": failures,
    }


def production_keyless_signing_readiness(
    policy_path: Path,
    signer_id: str,
    environment: dict[str, str] | None = None,
) -> dict[str, Any]:
    policy_result = validate_attestation_trust_policy(policy_path)
    failures: list[str] = []
    if policy_result.get("status") != "PASS":
        failures.append("attestation trust policy did not PASS validation")
    signer = policy_result.get("resolvedSigners", {}).get(signer_id)
    if not isinstance(signer, dict):
        failures.append(f"trusted signer id is absent from policy: {signer_id}")
        signer = {}
    if signer.get("mode") != "keyless":
        failures.append("production keyless signing readiness requires a keyless signer")
    if signer.get("certificateOidcIssuer") != GITHUB_ACTIONS_OIDC_ISSUER:
        failures.append("production keyless signing readiness currently admits only GitHub Actions OIDC")
    github = signer.get("githubActions")
    if not isinstance(github, dict):
        failures.append("production GitHub Actions keyless signer lacks exact workflow claims")
        github = {}
    env = dict(os.environ if environment is None else environment)
    sha_authority = github.get("shaAuthority")
    expected_sha = github.get("sha") if sha_authority == "policy" else env.get("ORDIVON_RELEASE_WORKFLOW_SHA")
    if sha_authority == "runtime":
        if not isinstance(expected_sha, str) or not re.fullmatch(r"[0-9a-f]{40}", expected_sha):
            failures.append("protected runtime did not supply ORDIVON_RELEASE_WORKFLOW_SHA as exact 40-hex occurrence authority")
    expected_runtime = {
        "GITHUB_REPOSITORY": github.get("repository"),
        "GITHUB_WORKFLOW_REF": github.get("workflowRef"),
        "GITHUB_REF": github.get("ref"),
        "GITHUB_SHA": expected_sha,
        "GITHUB_WORKFLOW": github.get("name"),
        "GITHUB_EVENT_NAME": github.get("trigger"),
    }
    if env.get("GITHUB_ACTIONS") != "true":
        failures.append("runtime is not an admitted GitHub Actions job")
    for key, expected in expected_runtime.items():
        if not isinstance(expected, str) or not expected:
            failures.append(f"keyless policy lacks expected runtime claim: {key}")
        elif env.get(key) != expected:
            failures.append(f"GitHub Actions runtime claim mismatch: {key}")
    token_url_present = bool(env.get("ACTIONS_ID_TOKEN_REQUEST_URL"))
    token_request_authority_present = bool(env.get("ACTIONS_ID_TOKEN_REQUEST_TOKEN"))
    if not token_url_present:
        failures.append("GitHub Actions OIDC request URL is unavailable; id-token: write is not usable")
    if not token_request_authority_present:
        failures.append("GitHub Actions OIDC request authority is unavailable; id-token: write is not usable")
    return {
        "status": "PASS" if not failures else "FAIL",
        "signerId": signer_id,
        "certificateIdentity": signer.get("certificateIdentity"),
        "certificateOidcIssuer": signer.get("certificateOidcIssuer"),
        "githubActions": github,
        "runtimeClaims": {key: env.get(key) for key in expected_runtime},
        "oidcRequestUrlPresent": token_url_present,
        "oidcRequestAuthorityPresent": token_request_authority_present,
        "failures": failures,
        "boundary": "PASS is an issuance preflight only: the runtime is an exact policy-bound GitHub Actions occurrence with OIDC request authority available. It does not claim Fulcio certificate issuance, Rekor logging, signing, or release completion until those external effects are separately observed.",
    }


def _cosign_executable() -> Path | None:
    configured = os.environ.get("ARTIFACT_COSIGN")
    if configured:
        path = Path(configured)
        return path if path.is_file() and os.access(path, os.X_OK) else None
    if COSIGN_SELECTED_BINARY.is_file() and os.access(COSIGN_SELECTED_BINARY, os.X_OK):
        return COSIGN_SELECTED_BINARY
    candidate = shutil.which("cosign")
    return Path(candidate) if candidate else None


def _cosign_selection_provenance(
    executable: Path,
    executable_digest: str,
    cosign_lock: dict[str, Any],
) -> dict[str, Any]:
    failures: list[str] = []
    if cosign_lock.get("originStanding") != "ARCH_REPOSITORY_PACKAGE_SIGNATURE_VERIFIED":
        failures.append("Cosign origin standing is not an accepted signed Arch repository package")
    if not COSIGN_ARCH_PACKAGE.is_file():
        failures.append("signed Arch Cosign package is absent")
    if not COSIGN_ARCH_PACKAGE_SIGNATURE.is_file():
        failures.append("Arch Cosign package detached signature is absent")
    pacman_key = shutil.which("pacman-key")
    if not pacman_key:
        failures.append("pacman-key is unavailable for Cosign package provenance verification")
    bsdtar = shutil.which("bsdtar")
    if not bsdtar:
        failures.append("bsdtar is unavailable for signed Cosign package inspection")
    expected_binary_digest = cosign_lock.get("binarySha256")
    if not isinstance(expected_binary_digest, str) or executable_digest != expected_binary_digest:
        failures.append("selected Cosign binary digest does not match the locked signed-package binary digest")
    if failures:
        return {
            "status": "FAIL",
            "authority": "Arch Linux package signing keyring",
            "package": str(COSIGN_ARCH_PACKAGE),
            "signature": str(COSIGN_ARCH_PACKAGE_SIGNATURE),
            "failures": failures,
        }

    package_digest = sha256_file(COSIGN_ARCH_PACKAGE)
    signature_digest = sha256_file(COSIGN_ARCH_PACKAGE_SIGNATURE)
    cache_key = (package_digest, signature_digest, executable_digest)
    cached = _COSIGN_PROVENANCE_CACHE.get(cache_key)
    if cached is not None:
        return dict(cached)

    if package_digest != cosign_lock.get("archPackageSha256"):
        failures.append("Arch Cosign package digest does not match the lock")
    if signature_digest != cosign_lock.get("archPackageSignatureSha256"):
        failures.append("Arch Cosign package signature digest does not match the lock")

    signature_check = subprocess.run(
        [str(pacman_key), "--verify", str(COSIGN_ARCH_PACKAGE_SIGNATURE), str(COSIGN_ARCH_PACKAGE)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=30,
    )
    signature_text = signature_check.stdout + "\n" + signature_check.stderr
    expected_fingerprint = str(cosign_lock.get("archPackageSignerFingerprint", ""))
    if signature_check.returncode != 0:
        failures.append("Arch Cosign package signature verification failed")
    if not expected_fingerprint or expected_fingerprint not in signature_text:
        failures.append("Arch Cosign package signer fingerprint does not match the lock")
    if "Good signature" not in signature_text:
        failures.append("Arch Cosign package signature was not reported as good")

    pkginfo = subprocess.run(
        [str(bsdtar), "-xOf", str(COSIGN_ARCH_PACKAGE), ".PKGINFO"],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=30,
    )
    expected_package_version = str(cosign_lock.get("archPackageVersion", ""))
    if pkginfo.returncode != 0:
        failures.append("signed Arch Cosign package metadata could not be read")
    elif f"pkgver = {expected_package_version}" not in pkginfo.stdout:
        failures.append("signed Arch Cosign package version does not match the lock")

    embedded_digest: str | None = None
    extractor = subprocess.Popen(
        [str(bsdtar), "-xOf", str(COSIGN_ARCH_PACKAGE), "usr/bin/cosign"],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if extractor.stdout is not None:
        embedded_hash = hashlib.sha256()
        for chunk in iter(lambda: extractor.stdout.read(1024 * 1024), b""):
            embedded_hash.update(chunk)
        embedded_digest = embedded_hash.hexdigest()
        extractor.stdout.close()
    if extractor.stderr is not None:
        extractor_stderr = extractor.stderr.read().decode("utf-8", errors="replace")
        extractor.stderr.close()
    else:
        extractor_stderr = ""
    extractor_returncode = extractor.wait(timeout=30)
    if extractor_returncode != 0:
        failures.append(f"signed Arch Cosign package binary extraction failed: {extractor_stderr[-1000:]}")
    elif embedded_digest != expected_binary_digest:
        failures.append("selected Cosign binary is not byte-identical to usr/bin/cosign in the signed Arch package")

    result = {
        "status": "PASS" if not failures else "FAIL",
        "authority": "Arch Linux package signing keyring",
        "originStanding": cosign_lock.get("originStanding"),
        "package": {
            "path": str(COSIGN_ARCH_PACKAGE.resolve()),
            "version": expected_package_version,
            "sha256": package_digest,
        },
        "signature": {
            "path": str(COSIGN_ARCH_PACKAGE_SIGNATURE.resolve()),
            "sha256": signature_digest,
            "returnCode": signature_check.returncode,
            "signerFingerprint": expected_fingerprint,
            "signer": cosign_lock.get("archPackageSigner"),
            "status": cosign_lock.get("archPackageSignatureStatus"),
        },
        "embeddedBinarySha256": embedded_digest,
        "selectedBinarySha256": executable_digest,
        "failures": failures,
        "boundary": "PASS proves the selected verifier bytes are byte-identical to usr/bin/cosign inside the exact digest-pinned Arch package whose detached signature verifies under the local Arch package keyring and locked signer fingerprint. It does not claim the bytes are the upstream Sigstore release binary.",
    }
    _COSIGN_PROVENANCE_CACHE[cache_key] = dict(result)
    return result


def cosign_tool_fact() -> dict[str, Any]:
    executable = _cosign_executable()
    if executable is None:
        return {"status": "NOT_AVAILABLE", "error": "cosign executable is unavailable"}
    proc = subprocess.run(
        [str(executable), "version"],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=20,
    )
    output = proc.stdout + "\n" + proc.stderr
    match = re.search(r"GitVersion:\s*v(\d+)\.(\d+)\.(\d+)(?:[^\s]*)?", output)
    version = tuple(int(item) for item in match.groups()) if match else None
    digest = sha256_file(executable)
    lock_error: str | None = None
    cosign_lock: dict[str, Any] = {}
    if COSIGN_LOCK_PATH.is_file():
        try:
            value = load_json(COSIGN_LOCK_PATH)
            candidate = value.get("cosign", {}) if isinstance(value, dict) else {}
            if isinstance(candidate, dict):
                cosign_lock = candidate
            else:
                lock_error = "toolchain lock cosign entry is not an object"
        except Exception as error:
            lock_error = str(error)
    else:
        lock_error = "toolchain lock is absent"
    expected_digest = cosign_lock.get("binarySha256") if cosign_lock else None
    expected_version = cosign_lock.get("observedVersion") if cosign_lock else None
    digest_ok = isinstance(expected_digest, str) and digest == expected_digest
    version_text = ".".join(str(item) for item in version) if version else None
    version_ok = isinstance(expected_version, str) and version_text == expected_version
    provenance = (
        _cosign_selection_provenance(executable, digest, cosign_lock)
        if lock_error is None and cosign_lock
        else {"status": "FAIL", "failures": [lock_error or "Cosign lock is unavailable"]}
    )
    status_ok = (
        proc.returncode == 0
        and version is not None
        and digest_ok
        and version_ok
        and provenance.get("status") == "PASS"
    )
    return {
        "status": "PASS" if status_ok else "FAIL",
        "path": str(executable.resolve()),
        "sha256": digest,
        "lockedSha256": expected_digest,
        "lockedDigestMatched": digest_ok,
        "version": version_text,
        "lockedVersion": expected_version,
        "lockedVersionMatched": version_ok,
        "versionTuple": list(version) if version else None,
        "rawVersion": output.strip()[:4000],
        "exitCode": proc.returncode,
        "lockError": lock_error,
        "provenance": provenance,
    }


def _version_at_least(observed: Iterable[int], minimum: tuple[int, int, int]) -> bool:
    values = tuple(int(item) for item in observed)
    return values >= minimum


def verify_sigstore_attestation_bundle_shape(
    bundle_path: Path,
    statement_path: Path,
    accepted_media_types: Iterable[str],
    signer_mode: str,
) -> dict[str, Any]:
    failures: list[str] = []
    bundle = load_json(bundle_path)
    statement = load_json(statement_path)
    accepted = set(accepted_media_types)
    media_type = bundle.get("mediaType") if isinstance(bundle, dict) else None
    if media_type != SIGSTORE_BUNDLE_V03 or media_type not in accepted:
        failures.append("bundle is not an explicitly accepted standardized Sigstore v0.3 JSON bundle")
    envelope = bundle.get("dsseEnvelope") if isinstance(bundle, dict) else None
    signed_statement: Any = None
    if not isinstance(envelope, dict):
        failures.append("Sigstore bundle does not contain a DSSE envelope")
    else:
        if envelope.get("payloadType") != INTOTO_DSSE_PAYLOAD_TYPE:
            failures.append("DSSE payloadType is not application/vnd.in-toto+json")
        signatures = envelope.get("signatures")
        if not isinstance(signatures, list) or not signatures:
            failures.append("DSSE envelope contains no signature")
        payload = envelope.get("payload")
        if not isinstance(payload, str):
            failures.append("DSSE envelope payload is absent")
        else:
            try:
                decoded = base64.b64decode(payload, validate=True)
                signed_statement = json.loads(decoded)
            except Exception as error:
                failures.append(f"DSSE payload is not valid base64 JSON: {error}")
    if signed_statement is not None and signed_statement != statement:
        failures.append("detached attestation statement does not exactly match the JSON object authenticated by the DSSE bundle")
    verification_material = bundle.get("verificationMaterial") if isinstance(bundle, dict) else None
    tlog_entries: list[Any] = []
    if not isinstance(verification_material, dict):
        failures.append("Sigstore bundle verificationMaterial is absent")
    else:
        raw_entries = verification_material.get("tlogEntries")
        if isinstance(raw_entries, list):
            tlog_entries = raw_entries
        if signer_mode == "public-key" and not isinstance(verification_material.get("publicKey"), dict):
            failures.append("public-key signer requires standardized bundle publicKey verification material")
        if signer_mode == "keyless":
            if not isinstance(verification_material.get("certificate"), dict):
                failures.append("keyless signer requires standardized bundle leaf certificate verification material")
            if isinstance(verification_material.get("publicKey"), dict):
                failures.append("keyless signer bundle must not substitute raw publicKey verification material")
    return {
        "status": "PASS" if not failures else "FAIL",
        "bundle": file_fact(bundle_path),
        "mediaType": media_type,
        "payloadType": envelope.get("payloadType") if isinstance(envelope, dict) else None,
        "signatureCount": len(envelope.get("signatures", [])) if isinstance(envelope, dict) and isinstance(envelope.get("signatures"), list) else 0,
        "tlogEntryCount": len(tlog_entries),
        "signedStatementMatches": signed_statement == statement if signed_statement is not None else False,
        "failures": failures,
    }


def _cosign_verify_standard_bundle(
    bundle_path: Path,
    subject: Path,
    signer: dict[str, Any],
    tool: dict[str, Any],
    predicate_type: str,
) -> dict[str, Any]:
    failures: list[str] = []
    if tool.get("status") != "PASS":
        failures.append("cosign verifier provenance/version/digest validation failed")
    version_tuple = tuple(tool.get("versionTuple") or [])
    if version_tuple and not _version_at_least(version_tuple, COSIGN_STANDARD_BUNDLE_MIN_VERSION):
        failures.append(
            f"cosign {tool.get('version')} is below the minimum allowed for standardized bundle verification: "
            + ".".join(str(item) for item in COSIGN_STANDARD_BUNDLE_MIN_VERSION)
        )
    mode = str(signer.get("mode", ""))
    command: list[str] = []
    if not failures and tool.get("status") == "PASS":
        command = [
            str(tool["path"]),
            "verify-blob-attestation",
            "--bundle", str(bundle_path),
            "--check-claims=true",
            "--type", predicate_type,
        ]
        if mode == "public-key":
            key_path = signer.get("publicKeyPath")
            if not isinstance(key_path, str):
                failures.append("resolved trusted public key is unavailable")
            else:
                command.extend(["--key", key_path])
                if signer.get("requireTransparencyLog") is not True:
                    command.append("--insecure-ignore-tlog")
        elif mode == "keyless":
            command.extend([
                "--certificate-identity", str(signer.get("certificateIdentity")),
                "--certificate-oidc-issuer", str(signer.get("certificateOidcIssuer")),
            ])
            github = signer.get("githubActions")
            if isinstance(github, dict):
                command.extend([
                    "--certificate-github-workflow-repository", str(github.get("repository")),
                    "--certificate-github-workflow-ref", str(github.get("ref")),
                    "--certificate-github-workflow-name", str(github.get("name")),
                    "--certificate-github-workflow-trigger", str(github.get("trigger")),
                ])
                github_sha = github.get("sha")
                if isinstance(github_sha, str):
                    command.extend(["--certificate-github-workflow-sha", github_sha])
            trusted_root = signer.get("trustedRootPath")
            if isinstance(trusted_root, str):
                command.extend(["--trusted-root", trusted_root])
        else:
            failures.append(f"unsupported signer mode: {mode}")
    if failures:
        return {
            "status": "FAIL",
            "exitCode": None,
            "stdout": "",
            "stderr": "",
            "failures": failures,
            "commandPolicy": {
                "checkClaims": True,
                "predicateType": predicate_type,
                "transparencyLogRequired": signer.get("requireTransparencyLog") is True,
                "githubActions": signer.get("githubActions"),
            },
        }
    command.append(str(subject))
    proc = subprocess.run(
        command,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        timeout=60,
    )
    if proc.returncode != 0:
        failures.append("cosign cryptographic attestation verification failed")
    return {
        "status": "PASS" if not failures else "FAIL",
        "exitCode": proc.returncode,
        "stdout": proc.stdout.strip()[:4000],
        "stderr": proc.stderr.strip()[:4000],
        "failures": failures,
        "commandPolicy": {
            "checkClaims": True,
            "predicateType": predicate_type,
            "transparencyLogRequired": signer.get("requireTransparencyLog") is True,
            "githubActions": signer.get("githubActions"),
        },
    }


def verify_signed_verification_summary(
    statement_path: Path,
    bundle_path: Path,
    subject: Path,
    profile_path: Path,
    trust_policy_path: Path,
    signer_id: str,
) -> dict[str, Any]:
    failures: list[str] = []
    policy_result = validate_attestation_trust_policy(trust_policy_path)
    if policy_result.get("status") != "PASS":
        failures.append("attestation trust policy did not PASS validation")
    signer = policy_result.get("resolvedSigners", {}).get(signer_id)
    if not isinstance(signer, dict):
        failures.append(f"trusted signer id is absent from policy: {signer_id}")
        signer = {}
    semantic = verify_verification_summary(statement_path, subject, profile_path)
    if semantic.get("status") != "PASS":
        failures.append("VSA semantic/digest binding validation failed")
    verifier_id = semantic.get("verifier", {}).get("id") if isinstance(semantic.get("verifier"), dict) else None
    if signer and verifier_id not in set(signer.get("allowedVerifierIds", [])):
        failures.append("VSA verifier.id is not authorized for the selected signer")
    accepted = policy_result.get("policy", {}).get("acceptedBundleMediaTypes", []) if isinstance(policy_result.get("policy"), dict) else []
    mode = str(signer.get("mode", ""))
    bundle_shape = verify_sigstore_attestation_bundle_shape(bundle_path, statement_path, accepted, mode) if bundle_path.is_file() else {
        "status": "FAIL", "failures": ["Sigstore bundle is absent"]
    }
    if bundle_shape.get("status") != "PASS":
        failures.append("Sigstore bundle shape/statement binding validation failed")
    if signer.get("requireTransparencyLog") is True and bundle_shape.get("tlogEntryCount", 0) < 1:
        failures.append("trust policy requires transparency-log evidence but bundle has no tlog entry")

    tool = cosign_tool_fact()
    cosign_result: dict[str, Any] = {"status": "NOT_RUN", "failures": []}
    if not failures:
        cosign_result = _cosign_verify_standard_bundle(
            bundle_path, subject, signer, tool, SLSA_VERIFICATION_SUMMARY_V1
        )
        failures.extend(cosign_result.get("failures", []))
    elif tool.get("status") != "PASS":
        failures.append("cosign verifier provenance/version/digest validation failed")
    return {
        "status": "PASS" if not failures else "FAIL",
        "authenticity": "VERIFIED" if not failures else "NOT_VERIFIED",
        "signerId": signer_id,
        "signerMode": mode,
        "verifierId": verifier_id,
        "trustPolicy": policy_result,
        "bundleShape": bundle_shape,
        "semantic": semantic,
        "cosign": cosign_result,
        "tool": tool,
        "failures": failures,
        "boundary": "PASS requires a standardized Sigstore v0.3 DSSE bundle, exact signed-statement equality, an authorized signer->verifier.id mapping, exact VSA subject/policy/resource bindings, a provenance-verified digest/version-pinned Cosign verifier, and successful cosign verification with --check-claims=true. Keyless verification additionally requires identity/issuer policy plus transparency-log evidence. Legacy bundles are rejected before cosign verification.",
    }

def _write_raw_and_vsa(
    output_dir: Path,
    gate: str,
    subject: Path,
    profile_path: Path,
    raw: dict[str, Any],
    verifier_id: str,
    verifier_versions: dict[str, str],
    passed: bool | None = None,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    raw_path = output_dir / f"{gate}.raw.json"
    write_json(raw_path, raw)
    result_passed = raw.get("status") == "PASS" if passed is None else passed
    statement = verification_summary_statement(
        subject,
        profile_path,
        verifier_id,
        verifier_versions,
        result_passed,
    )
    statement_path = output_dir / f"{gate}.vsa.json"
    write_json(statement_path, statement)
    checked = verify_verification_summary(statement_path, subject, profile_path, [verifier_id])
    return {
        "gate": gate,
        "status": "PASS" if checked.get("status") == "PASS" and result_passed else "FAIL",
        "verificationResult": "PASSED" if result_passed else "FAILED",
        "rawEvidence": file_fact(raw_path),
        "vsa": file_fact(statement_path),
        "vsaValidation": checked,
    }


def execute_verify_stage(profile_path: Path, artifact: Path, output_dir: Path) -> dict[str, Any]:
    profile_result = validate_profile(profile_path)
    profile = profile_result.get("profile", {})
    artifact_class = profile.get("artifactClass")
    receipts: dict[str, Any] = {}
    failures: list[str] = []
    if profile_result.get("status") != "PASS":
        failures.append("delivery profile did not PASS validation")
    expected_suffix = _primary_suffix(profile) if profile else None
    if expected_suffix and artifact.suffix.casefold() != expected_suffix:
        failures.append(f"artifact suffix {artifact.suffix} does not match profile primary output {expected_suffix}")
    if not artifact.is_file():
        failures.append("artifact file is absent")
    if failures:
        return {
            "schemaVersion": 1,
            "kind": "artifact-delivery-verify-stage",
            "status": "FAIL",
            "profileId": profile.get("id"),
            "failures": failures,
            "receipts": receipts,
        }

    profile_raw = {
        "status": profile_result.get("status"),
        "profile": file_fact(profile_path),
        "jsonSchema": profile_result.get("jsonSchema"),
        "minimalContractErrors": profile_result.get("minimalContractErrors", []),
        "boundary": "Delivery-profile schema/contract validation only; artifact-format, target, visual, accessibility and delivery gates remain independent.",
    }
    receipts["profileSchema"] = _write_raw_and_vsa(
        output_dir,
        "profileSchema",
        artifact,
        profile_path,
        profile_raw,
        LOCAL_VSA_VERIFIER_ID,
        {"jsonschema": importlib.metadata.version("jsonschema")},
    )

    if artifact_class in {"presentation", "document", "spreadsheet"}:
        raw = verify_openxml_artifact(artifact)
        version = str(raw.get("validatorOutput", {}).get("validator", {}).get("packageVersion", "unknown"))
        receipts["structural"] = _write_raw_and_vsa(
            output_dir,
            "structural",
            artifact,
            profile_path,
            raw,
            LOCAL_VSA_VERIFIER_ID,
            {"DocumentFormat.OpenXml": version},
        )
        if artifact_class == "presentation":
            inspected = inspect_pptx(artifact, profile.get("semanticPolicy", {}).get("placeholderPatterns", []))
            semantic = verify_presentation_semantics(profile, inspected)
            raw_semantic = {
                "status": "PASS" if inspected.get("status") == "PASS" and semantic.get("status") == "PASS" else "FAIL",
                "artifact": file_fact(artifact),
                "inspection": inspected,
                "semantic": semantic,
                "boundary": "Presentation-local package/slide semantic checks only; target rendering, visual review and delivery remain separate.",
            }
            receipts["semantic"] = _write_raw_and_vsa(
                output_dir,
                "semantic",
                artifact,
                profile_path,
                raw_semantic,
                LOCAL_VSA_VERIFIER_ID,
                {"python-pptx": importlib.metadata.version("python-pptx")},
            )
    elif artifact_class in {"fixed-view", "archive", "accessible"}:
        raw = verify_pdf(artifact)
        qpdf_version = "unknown"
        if shutil.which("qpdf"):
            qproc = subprocess.run([shutil.which("qpdf"), "--version"], text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False, timeout=10)
            qpdf_version = qproc.stdout.splitlines()[0] if qproc.stdout else "unknown"
        receipts["structural"] = _write_raw_and_vsa(
            output_dir,
            "structural",
            artifact,
            profile_path,
            raw,
            LOCAL_VSA_VERIFIER_ID,
            {"qpdf": qpdf_version},
        )
        if profile.get("gates", {}).get("conformance") is True:
            flavour = profile.get("conformancePolicy", {}).get("pdfFlavour")
            if not isinstance(flavour, str):
                conformance_raw = {"status": "FAIL", "error": "profile requires conformance but omits conformancePolicy.pdfFlavour"}
            else:
                conformance_raw = verify_pdf_conformance(artifact, flavour)
            receipts["conformance"] = _write_raw_and_vsa(
                output_dir,
                "conformance",
                artifact,
                profile_path,
                conformance_raw,
                LOCAL_VSA_VERIFIER_ID,
                {"veraPDF": "1.30.2"},
            )
    elif artifact_class == "web":
        conformance = verify_html_conformance(artifact)
        receipts["conformance"] = _write_raw_and_vsa(
            output_dir,
            "conformance",
            artifact,
            profile_path,
            conformance,
            LOCAL_VSA_VERIFIER_ID,
            {"Nu Html Checker": str(conformance.get("validator", {}).get("version") or "unknown")},
        )
        receipts["structural"] = _write_raw_and_vsa(
            output_dir,
            "structural",
            artifact,
            profile_path,
            conformance,
            LOCAL_VSA_VERIFIER_ID,
            {"Nu Html Checker": str(conformance.get("validator", {}).get("version") or "unknown")},
        )
        web = verify_web_local(artifact)
        output = web.get("verifierOutput", {}) if isinstance(web.get("verifierOutput"), dict) else {}
        browsers = output.get("browsers", {}) if isinstance(output.get("browsers"), dict) else {}
        chromium = browsers.get("chromium", {}) if isinstance(browsers.get("chromium"), dict) else {}
        accessibility_passed = chromium.get("status") == "PASS" and chromium.get("accessibility", {}).get("status") == "PASS"
        receipts["accessibility"] = _write_raw_and_vsa(
            output_dir,
            "accessibility",
            artifact,
            profile_path,
            web,
            LOCAL_VSA_VERIFIER_ID,
            {
                "@axe-core/playwright": str(output.get("tooling", {}).get("axePlaywright", "unknown")),
                "@playwright/test": str(output.get("tooling", {}).get("playwright", "unknown")),
            },
            accessibility_passed,
        )
        required_names = [str(profile.get("targetRenderer", {}).get("name"))] if profile.get("targetRenderer", {}).get("required") is True else []
        required_names.extend(
            str(item.get("name"))
            for item in profile.get("secondaryRenderers", [])
            if item.get("required") is True
        )
        browser_by_name = {
            str(item.get("name")): item
            for item in browsers.values()
            if isinstance(item, dict) and item.get("name")
        }
        target_passed = bool(required_names) and all(browser_by_name.get(name, {}).get("status") == "PASS" for name in required_names)
        target_raw = {
            "status": "PASS" if target_passed else "FAIL",
            "artifact": file_fact(artifact),
            "requiredRenderers": required_names,
            "browserResults": browser_by_name,
            "failures": [name for name in required_names if browser_by_name.get(name, {}).get("status") != "PASS"],
            "boundary": "Target renderer policy requires every profile-required primary/secondary renderer. Unsupported-host WebKit remains a hard failure rather than a local compatibility waiver.",
        }
        receipts["target"] = _write_raw_and_vsa(
            output_dir,
            "target",
            artifact,
            profile_path,
            target_raw,
            LOCAL_VSA_VERIFIER_ID,
            {"@playwright/test": str(output.get("tooling", {}).get("playwright", "unknown"))},
            target_passed,
        )
    else:
        failures.append(f"no verify-stage adapters for artifactClass={artifact_class!r}")

    executed_failures = [gate for gate, item in receipts.items() if item.get("status") != "PASS"]
    required_gates = {name for name, flag in profile.get("gates", {}).items() if flag is True}
    generated = set(receipts)
    pending = sorted(required_gates - generated)
    if executed_failures:
        failures.append("executed verifier gate(s) failed: " + ", ".join(sorted(executed_failures)))
    return {
        "schemaVersion": 1,
        "kind": "artifact-delivery-verify-stage",
        "status": "PASS" if not failures else "FAIL",
        "profileId": profile.get("id"),
        "artifact": file_fact(artifact),
        "receipts": receipts,
        "profileRequiredGates": sorted(required_gates),
        "pendingRequiredGates": pending,
        "profileVerificationComplete": not failures and not pending,
        "failures": failures,
        "boundary": "Verify-stage status covers only adapters executed here. profileVerificationComplete is the stronger statement and remains false until every profile-required gate has independent evidence. Generated VSAs are unsigned local statements; external trust requires signature/root-of-trust verification.",
    }


def aggregate_vsa_gates(
    profile_path: Path,
    subject: Path,
    gate_paths: dict[str, Path],
    allow_local_unsigned: bool = False,
    bundles: dict[str, Path] | None = None,
    trust_policy_path: Path | None = None,
    signer_ids: dict[str, str] | None = None,
) -> dict[str, Any]:
    profile_result = validate_profile(profile_path)
    profile = profile_result.get("profile", {})
    required_all = {name for name, flag in profile.get("gates", {}).items() if flag is True}
    required = required_all & VSA_GATE_NAMES
    assembly_required = required_all & ASSEMBLY_GATE_NAMES
    components: dict[str, Any] = {}
    failures: list[str] = []
    bundle_paths = bundles or {}
    selected_signers = signer_ids or {}
    for gate in sorted(set(bundle_paths) - set(gate_paths)):
        failures.append(f"Sigstore bundle supplied without a corresponding VSA gate: {gate}")
    for gate in sorted(set(selected_signers) - set(gate_paths)):
        failures.append(f"trusted signer supplied without a corresponding VSA gate: {gate}")
    for gate, path in sorted(gate_paths.items()):
        if gate not in profile.get("gates", {}):
            failures.append(f"VSA supplied for undeclared gate: {gate}")
            continue
        if gate not in VSA_GATE_NAMES:
            failures.append(f"gate is assembly/provenance policy and must not be represented as a VSA gate: {gate}")
            continue
        checked = verify_verification_summary(path, subject, profile_path)
        verification_result = checked.get("verificationResult")
        signature_result: dict[str, Any] | None = None
        if checked.get("status") != "PASS":
            failures.append(f"VSA binding/shape validation failed: {gate}")
        elif verification_result != "PASSED":
            failures.append(f"VSA verificationResult is not PASSED: {gate}")
        if not allow_local_unsigned:
            bundle_path = bundle_paths.get(gate)
            signer_id = selected_signers.get(gate)
            if trust_policy_path is None:
                failures.append(f"VSA authenticity trust policy missing for production gate: {gate}")
            elif bundle_path is None:
                failures.append(f"VSA authenticity Sigstore bundle missing for production gate: {gate}")
            elif not signer_id:
                failures.append(f"VSA authenticity trusted signer id missing for production gate: {gate}")
            else:
                signature_result = verify_signed_verification_summary(
                    path, bundle_path, subject, profile_path, trust_policy_path, signer_id
                )
                if signature_result.get("status") != "PASS":
                    failures.append(f"VSA authenticity verification failed: {gate}")
        component = dict(checked)
        component["signature"] = signature_result
        component["authenticity"] = (
            "LOCAL_UNSIGNED_DEVELOPMENT" if allow_local_unsigned
            else "VERIFIED" if signature_result and signature_result.get("status") == "PASS"
            else "NOT_VERIFIED"
        )
        components[gate] = component
    for gate in sorted(required):
        if gate not in components:
            failures.append(f"required gate VSA missing: {gate}")
    if profile_result.get("status") != "PASS":
        failures.append("profile schema did not PASS")
    return {
        "schemaVersion": 1,
        "kind": "artifact-delivery-vsa-gate-aggregation",
        "status": "PASS" if not failures else "FAIL",
        "profileId": profile.get("id"),
        "subject": file_fact(subject),
        "requiredGates": sorted(required),
        "assemblyGates": sorted(assembly_required),
        "components": components,
        "allowLocalUnsigned": allow_local_unsigned,
        "trustPolicy": file_fact(trust_policy_path) if trust_policy_path is not None and trust_policy_path.is_file() else None,
        "bundles": {gate: file_fact(path) for gate, path in sorted(bundle_paths.items()) if path.is_file()},
        "selectedSigners": dict(sorted(selected_signers.items())),
        "failures": failures,
        "boundary": "Production aggregation requires a standardized Sigstore bundle plus a policy-authorized signer for every required VSA gate. Local unsigned mode is only for same-workspace development evidence and must not be presented as external trust.",
    }


def verify_file_fact(fact: dict[str, Any]) -> dict[str, Any]:
    failures: list[str] = []
    path_value = fact.get("path")
    path = Path(path_value) if isinstance(path_value, str) and path_value else None
    if path is None or not path.is_file():
        failures.append("referenced file is absent")
        return {"status": "FAIL", "path": path_value, "failures": failures}
    actual = file_fact(path)
    if fact.get("name") != actual.get("name"):
        failures.append("file name mismatch")
    if fact.get("size") != actual.get("size"):
        failures.append("file size mismatch")
    if fact.get("digest", {}).get("sha256") != actual.get("digest", {}).get("sha256"):
        failures.append("file SHA-256 mismatch")
    return {"status": "PASS" if not failures else "FAIL", "path": str(path), "actual": actual, "failures": failures}


def _copy_exact(
    source: Path,
    destination: Path,
    expected_source_fact: dict[str, Any] | None = None,
) -> dict[str, Any]:
    source_before = file_fact(source)
    expected_digest = (
        expected_source_fact.get("digest", {}).get("sha256")
        if isinstance(expected_source_fact, dict)
        else source_before["digest"]["sha256"]
    )
    expected_size = expected_source_fact.get("size") if isinstance(expected_source_fact, dict) else source_before["size"]
    expected_name = expected_source_fact.get("name") if isinstance(expected_source_fact, dict) else source_before["name"]
    expected_matched = (
        source_before["digest"]["sha256"] == expected_digest
        and source_before["size"] == expected_size
        and source_before["name"] == expected_name
    )
    if not expected_matched:
        return {
            "status": "FAIL",
            "source": source_before,
            "destination": None,
            "expectedSource": expected_source_fact,
            "expectedSourceMatched": False,
            "sourceStable": False,
            "digestMatched": False,
        }
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, destination)
    source_after = file_fact(source)
    destination_fact = file_fact(destination)
    source_stable = (
        source_after["digest"]["sha256"] == source_before["digest"]["sha256"]
        and source_after["size"] == source_before["size"]
    )
    destination_matched = (
        destination_fact["digest"]["sha256"] == expected_digest
        and destination_fact["size"] == expected_size
    )
    return {
        "status": "PASS" if expected_matched and source_stable and destination_matched else "FAIL",
        "source": source_after,
        "destination": destination_fact,
        "expectedSource": expected_source_fact,
        "expectedSourceMatched": expected_matched,
        "sourceStable": source_stable,
        "digestMatched": destination_matched,
    }


def _relative_file_fact(root: Path, path: Path) -> dict[str, Any]:
    fact = file_fact(path)
    fact["path"] = path.resolve().relative_to(root.resolve()).as_posix()
    return fact


def execute_package_stage(
    profile_path: Path,
    primary: Path,
    verify_report_path: Path,
    output_dir: Path,
    companions: Iterable[Path] = (),
    request_path: Path | None = None,
    provenance_path: Path | None = None,
    provenance_bundle_path: Path | None = None,
    provenance_signer_id: str | None = None,
    allow_local_unsigned: bool = False,
    gate_bundles: dict[str, Path] | None = None,
    trust_policy_path: Path | None = None,
    signer_ids: dict[str, str] | None = None,
) -> dict[str, Any]:
    profile_result = validate_profile(profile_path)
    profile = profile_result.get("profile", {})
    failures: list[str] = []
    companion_paths = list(companions)
    if profile_result.get("status") != "PASS":
        failures.append("delivery profile did not PASS validation")
    if not primary.is_file():
        failures.append("primary artifact is absent")
    verify_report_fact = file_fact(verify_report_path)
    verify_report = load_json(verify_report_path)
    if verify_report.get("kind") != "artifact-delivery-verify-stage":
        failures.append("verify report kind is not artifact-delivery-verify-stage")
    report_artifact = verify_report.get("artifact", {}) if isinstance(verify_report.get("artifact"), dict) else {}
    if report_artifact.get("digest", {}).get("sha256") != sha256_file(primary):
        failures.append("verify report artifact digest does not bind the primary artifact")
    if report_artifact.get("name") != primary.name:
        failures.append("verify report artifact name does not bind the primary artifact")

    receipts = verify_report.get("receipts", {}) if isinstance(verify_report.get("receipts"), dict) else {}
    gate_paths: dict[str, Path] = {}
    raw_paths: dict[str, Path] = {}
    expected_raw_facts: dict[str, dict[str, Any]] = {}
    expected_vsa_facts: dict[str, dict[str, Any]] = {}
    receipt_checks: dict[str, Any] = {}
    for gate, receipt in sorted(receipts.items()):
        if not isinstance(receipt, dict):
            failures.append(f"invalid verify receipt object: {gate}")
            continue
        raw_fact = receipt.get("rawEvidence", {}) if isinstance(receipt.get("rawEvidence"), dict) else {}
        vsa_fact = receipt.get("vsa", {}) if isinstance(receipt.get("vsa"), dict) else {}
        raw_check = verify_file_fact(raw_fact)
        vsa_check = verify_file_fact(vsa_fact)
        receipt_checks[gate] = {"raw": raw_check, "vsa": vsa_check}
        if raw_check.get("status") != "PASS":
            failures.append(f"raw evidence file fact failed: {gate}")
        if vsa_check.get("status") != "PASS":
            failures.append(f"VSA file fact failed: {gate}")
        if raw_check.get("status") == "PASS":
            raw_paths[gate] = Path(raw_check["path"])
            expected_raw_facts[gate] = raw_fact
        if vsa_check.get("status") == "PASS":
            gate_paths[gate] = Path(vsa_check["path"])
            expected_vsa_facts[gate] = vsa_fact

    bundle_paths = gate_bundles or {}
    selected_signers = signer_ids or {}
    trust_policy_fact = (
        file_fact(trust_policy_path)
        if trust_policy_path is not None and trust_policy_path.is_file()
        else None
    )
    gate_aggregation = aggregate_vsa_gates(
        profile_path,
        primary,
        gate_paths,
        allow_local_unsigned,
        bundle_paths,
        trust_policy_path,
        selected_signers,
    )
    if gate_aggregation.get("status") != "PASS":
        failures.append("required VSA gate aggregation did not PASS")

    provenance_required = profile.get("gates", {}).get("releaseProvenance") is True
    provenance_preflight: dict[str, Any] | None = None
    expected_provenance_fact: dict[str, Any] | None = None
    expected_provenance_bundle_fact: dict[str, Any] | None = None
    if provenance_path is not None:
        if not provenance_path.is_file():
            failures.append("release provenance statement is absent")
        elif allow_local_unsigned:
            semantic = verify_release_provenance(
                provenance_path,
                [primary, *companion_paths],
                profile_path=profile_path,
                request_path=request_path,
            )
            provenance_preflight = {
                "status": semantic.get("status"),
                "authenticity": "LOCAL_UNSIGNED_DEVELOPMENT",
                "semantic": semantic,
                "failures": list(semantic.get("failures", [])),
            }
            expected_provenance_fact = file_fact(provenance_path)
            if semantic.get("status") != "PASS":
                failures.append("release provenance semantic/request binding did not PASS")
        else:
            if request_path is None:
                failures.append("production release provenance verification requires the digest-bound delivery request")
            if trust_policy_path is None:
                failures.append("production release provenance verification requires an attestation trust policy")
            if provenance_bundle_path is None:
                failures.append("production release provenance verification requires a Sigstore bundle")
            if not provenance_signer_id:
                failures.append("production release provenance verification requires a trusted signer id")
            if (
                request_path is not None
                and trust_policy_path is not None
                and provenance_bundle_path is not None
                and provenance_signer_id
            ):
                provenance_preflight = verify_signed_release_provenance(
                    provenance_path,
                    provenance_bundle_path,
                    [primary, *companion_paths],
                    profile_path,
                    request_path,
                    trust_policy_path,
                    provenance_signer_id,
                )
                expected_provenance_fact = file_fact(provenance_path)
                expected_provenance_bundle_fact = file_fact(provenance_bundle_path) if provenance_bundle_path.is_file() else None
                if provenance_preflight.get("status") != "PASS":
                    failures.append("release provenance cryptographic authenticity verification failed")
    elif provenance_required:
        if allow_local_unsigned:
            if request_path is None:
                failures.append("development release provenance generation requires the digest-bound delivery request")
        else:
            failures.append("production release provenance requires a pre-existing signed SLSA Provenance statement")

    if failures:
        return {
            "schemaVersion": 1,
            "kind": "artifact-delivery-package-stage",
            "status": "FAIL",
            "profileId": profile.get("id"),
            "primary": file_fact(primary) if primary.is_file() else None,
            "verifyReport": file_fact(verify_report_path),
            "receiptChecks": receipt_checks,
            "gateAggregation": gate_aggregation,
            "provenancePreflight": provenance_preflight,
            "packageCreated": False,
            "releaseReady": False,
            "failures": failures,
            "boundary": "Package-stage preflight failed before copying release bytes. Missing/failed/authenticity-unverified required gates cannot be converted into package PASS.",
        }

    if output_dir.exists() and any(output_dir.iterdir()):
        return {
            "schemaVersion": 1,
            "kind": "artifact-delivery-package-stage",
            "status": "FAIL",
            "profileId": profile.get("id"),
            "packageCreated": False,
            "releaseReady": False,
            "failures": ["package output directory already exists and is not empty"],
        }
    output_dir.mkdir(parents=True, exist_ok=True)
    artifacts_dir = output_dir / "artifacts"
    evidence_dir = output_dir / "evidence"
    attestations_dir = output_dir / "attestations"
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    attestations_dir.mkdir(parents=True, exist_ok=True)

    copy_results: list[dict[str, Any]] = []
    package_primary = artifacts_dir / primary.name
    copy_results.append(_copy_exact(primary, package_primary, report_artifact))
    package_companions: list[Path] = []
    seen_artifact_names = {primary.name}
    for companion in companion_paths:
        if companion.name in seen_artifact_names:
            failures.append(f"artifact basename collision: {companion.name}")
            continue
        seen_artifact_names.add(companion.name)
        destination = artifacts_dir / companion.name
        copy_results.append(_copy_exact(companion, destination))
        package_companions.append(destination)

    package_evidence: list[Path] = []
    package_vsas: list[Path] = []
    package_sigstore_bundles: list[Path] = []
    package_gate_paths: dict[str, Path] = {}
    package_bundle_paths: dict[str, Path] = {}
    for gate in sorted(gate_paths):
        raw_destination = evidence_dir / f"{gate}.raw.json"
        vsa_destination = attestations_dir / f"{gate}.vsa.json"
        copy_results.append(_copy_exact(raw_paths[gate], raw_destination, expected_raw_facts[gate]))
        copy_results.append(_copy_exact(gate_paths[gate], vsa_destination, expected_vsa_facts[gate]))
        package_evidence.append(raw_destination)
        package_vsas.append(vsa_destination)
        package_gate_paths[gate] = vsa_destination
        if gate in bundle_paths:
            bundle_destination = attestations_dir / f"{gate}.sigstore.json"
            expected_bundle_fact = gate_aggregation.get("bundles", {}).get(gate)
            copy_results.append(_copy_exact(bundle_paths[gate], bundle_destination, expected_bundle_fact))
            package_sigstore_bundles.append(bundle_destination)
            package_bundle_paths[gate] = bundle_destination
    verify_destination = evidence_dir / "verify-stage.json"
    copy_results.append(_copy_exact(verify_report_path, verify_destination, verify_report_fact))
    package_evidence.append(verify_destination)

    if any(item.get("status") != "PASS" for item in copy_results):
        failures.append("one or more exact package copies failed expected-digest/stability comparison")

    if trust_policy_fact is not None and verify_file_fact(trust_policy_fact).get("status") != "PASS":
        failures.append("attestation trust policy changed after preflight verification")
    packaged_gate_aggregation = aggregate_vsa_gates(
        profile_path,
        package_primary,
        package_gate_paths,
        allow_local_unsigned,
        package_bundle_paths,
        trust_policy_path,
        selected_signers,
    )
    if packaged_gate_aggregation.get("status") != "PASS":
        failures.append("packaged VSA gate aggregation did not PASS re-verification")
    gate_aggregation_path = evidence_dir / "vsa-gate-aggregation.json"
    write_json(gate_aggregation_path, packaged_gate_aggregation)
    package_evidence.append(gate_aggregation_path)

    provenance_destination: Path | None = None
    provenance_bundle_destination: Path | None = None
    packaged_provenance_verification: dict[str, Any] | None = None
    if provenance_path is not None:
        provenance_destination = attestations_dir / "slsa-provenance.json"
        copy_results.append(_copy_exact(provenance_path, provenance_destination, expected_provenance_fact))
        if provenance_bundle_path is not None:
            provenance_bundle_destination = attestations_dir / "slsa-provenance.sigstore.json"
            copy_results.append(_copy_exact(
                provenance_bundle_path,
                provenance_bundle_destination,
                expected_provenance_bundle_fact,
            ))
    elif provenance_required and allow_local_unsigned and request_path is not None:
        request_validation = validate_delivery_request(request_path)
        if request_validation.get("status") != "PASS":
            failures.append("delivery request did not PASS validation for provenance generation")
        else:
            request = request_validation["request"]
            profile_ref = request.get("profile", {})
            if profile_ref.get("id") != profile.get("id") or profile_ref.get("sha256") != sha256_file(profile_path):
                failures.append("delivery request profile binding does not match package profile")
            else:
                materials = [Path(request_validation["resolved"]["source"]["path"])]
                materials.extend(Path(item["path"]) for item in request_validation["resolved"].get("materials", []))
                provenance_destination = attestations_dir / "slsa-provenance.json"
                request_plan = compile_delivery_plan(request_path)
                if request_plan.get("status") != "PASS":
                    failures.append("delivery request did not compile to an accepted build definition for local provenance generation")
                else:
                    provenance = slsa_statement(
                        [package_primary, *package_companions],
                        materials,
                        profile_path,
                        LOCAL_BUILD_PLATFORM_ID,
                        str(request_plan["buildType"]),
                        request_path=request_path,
                    )
                    write_json(provenance_destination, provenance)

    if (
        any(item.get("status") != "PASS" for item in copy_results)
        and "one or more exact package copies failed expected-digest/stability comparison" not in failures
    ):
        failures.append("one or more exact package copies failed expected-digest/stability comparison")

    if provenance_destination is not None:
        if allow_local_unsigned:
            semantic = verify_release_provenance(
                provenance_destination,
                [package_primary, *package_companions],
                profile_path=profile_path,
                request_path=request_path,
            )
            packaged_provenance_verification = {
                "status": semantic.get("status"),
                "authenticity": "LOCAL_UNSIGNED_DEVELOPMENT",
                "semantic": semantic,
                "failures": list(semantic.get("failures", [])),
            }
        elif (
            provenance_bundle_destination is not None
            and request_path is not None
            and trust_policy_path is not None
            and provenance_signer_id
        ):
            packaged_provenance_verification = verify_signed_release_provenance(
                provenance_destination,
                provenance_bundle_destination,
                [package_primary, *package_companions],
                profile_path,
                request_path,
                trust_policy_path,
                provenance_signer_id,
            )
        else:
            packaged_provenance_verification = {
                "status": "FAIL",
                "authenticity": "NOT_VERIFIED",
                "failures": ["packaged provenance lacks required production authenticity inputs"],
            }
        if packaged_provenance_verification.get("status") != "PASS":
            failures.append("packaged release provenance did not PASS re-verification")

    release_manifest = build_release_manifest(
        profile_path,
        package_primary,
        package_companions,
        [
            *package_evidence,
            *package_vsas,
            *package_sigstore_bundles,
            *([provenance_bundle_destination] if provenance_bundle_destination is not None else []),
        ],
        provenance_destination,
    )
    if release_manifest.get("status") != "PASS":
        failures.append("release manifest assembly did not PASS")
    release_manifest_path = output_dir / "release-manifest.json"
    write_json(release_manifest_path, release_manifest)

    package_status = "PASS" if not failures else "FAIL"
    package_index = {
        "schemaVersion": 1,
        "kind": "artifact-delivery-package-index",
        "status": package_status,
        "profile": {"id": profile.get("id"), "sha256": sha256_file(profile_path)},
        "primary": _relative_file_fact(output_dir, package_primary),
        "companions": [_relative_file_fact(output_dir, path) for path in package_companions],
        "evidence": [_relative_file_fact(output_dir, path) for path in package_evidence],
        "attestations": [
            *[_relative_file_fact(output_dir, path) for path in package_vsas],
            *[_relative_file_fact(output_dir, path) for path in package_sigstore_bundles],
            *([_relative_file_fact(output_dir, provenance_destination)] if provenance_destination and provenance_destination.is_file() else []),
            *([_relative_file_fact(output_dir, provenance_bundle_destination)] if provenance_bundle_destination and provenance_bundle_destination.is_file() else []),
        ],
        "trustPolicy": (
            {
                "id": load_json(trust_policy_path).get("id"),
                "sha256": sha256_file(trust_policy_path),
            }
            if trust_policy_path is not None and trust_policy_path.is_file()
            else None
        ),
        "releaseManifest": _relative_file_fact(output_dir, release_manifest_path),
        "provenanceTrust": packaged_provenance_verification,
        "trustStanding": "LOCAL_UNSIGNED_DEVELOPMENT" if allow_local_unsigned else "CRYPTOGRAPHICALLY_VERIFIED",
        "releaseReady": package_status == "PASS" and not allow_local_unsigned,
    }
    package_index_path = output_dir / "package-index.json"
    write_json(package_index_path, package_index)
    return {
        "schemaVersion": 1,
        "kind": "artifact-delivery-package-stage",
        "status": package_status,
        "profileId": profile.get("id"),
        "packageDirectory": str(output_dir.resolve()),
        "packageCreated": package_status == "PASS",
        "packageIndex": file_fact(package_index_path),
        "releaseManifest": release_manifest,
        "preflightGateAggregation": gate_aggregation,
        "gateAggregation": packaged_gate_aggregation,
        "preflightProvenance": provenance_preflight,
        "provenanceVerification": packaged_provenance_verification,
        "copyResults": copy_results,
        "trustStanding": package_index["trustStanding"],
        "releaseReady": package_index["releaseReady"],
        "failures": failures,
        "boundary": "Package PASS means exact bytes/evidence/attestations were assembled and the release manifest passed. Production releaseReady additionally requires every required VSA gate and every supplied/profile-required SLSA Provenance statement to pass policy-bound cryptographic Sigstore verification before and after exact package copy; local unsigned mode remains development-only and explicitly not release-ready.",
    }


def _slsa_resource_descriptor(path: Path) -> dict[str, Any]:
    digest = sha256_file(path)
    return {
        "uri": ni_sha256_uri(path),
        "name": path.name,
        "digest": {"sha256": digest},
    }


def verify_release_provenance(
    provenance_path: Path,
    subjects: Iterable[Path],
    profile_path: Path | None = None,
    request_path: Path | None = None,
) -> dict[str, Any]:
    failures: list[str] = []
    value = load_json(provenance_path)
    subject_paths = list(subjects)
    expected = {path.name: sha256_file(path) for path in subject_paths}
    if value.get("_type") != IN_TOTO_STATEMENT_V1:
        failures.append("provenance is not an in-toto Statement v1")
    if value.get("predicateType") != SLSA_PROVENANCE_V1:
        failures.append("provenance predicateType is not SLSA Provenance v1")
    observed: dict[str, str] = {}
    for item in value.get("subject", []) if isinstance(value.get("subject"), list) else []:
        if isinstance(item, dict) and isinstance(item.get("name"), str):
            digest = item.get("digest", {}).get("sha256") if isinstance(item.get("digest"), dict) else None
            if isinstance(digest, str):
                observed[item["name"]] = digest
    if observed != expected:
        failures.append("provenance subjects do not exactly bind the release primary/companions")

    predicate = value.get("predicate", {}) if isinstance(value.get("predicate"), dict) else {}
    build_definition = predicate.get("buildDefinition", {}) if isinstance(predicate.get("buildDefinition"), dict) else {}
    run_details = predicate.get("runDetails", {}) if isinstance(predicate.get("runDetails"), dict) else {}
    builder = run_details.get("builder", {}).get("id") if isinstance(run_details.get("builder"), dict) else None
    build_type = build_definition.get("buildType")
    external_parameters = build_definition.get("externalParameters")
    resolved_dependencies = build_definition.get("resolvedDependencies")
    if not isinstance(builder, str) or not urllib.parse.urlparse(builder).scheme:
        failures.append("provenance builder.id is absent or not a URI")
    if not isinstance(build_type, str) or not urllib.parse.urlparse(build_type).scheme:
        failures.append("provenance buildType is absent or not a URI")
    if not isinstance(external_parameters, dict):
        failures.append("provenance externalParameters is absent or not an object")
        external_parameters = {}
    if not isinstance(resolved_dependencies, list):
        failures.append("provenance resolvedDependencies is absent or not an array")
        resolved_dependencies = []

    expected_external: dict[str, Any] = {}
    expected_dependencies: list[dict[str, Any]] = []
    expected_build_type: str | None = None
    request_validation: dict[str, Any] | None = None
    if profile_path is not None:
        profile = load_json(profile_path)
        expected_external["deliveryProfile"] = {
            "id": profile.get("id"),
            "sha256": sha256_file(profile_path),
        }
        expected_dependencies.append(_slsa_resource_descriptor(profile_path))
    if request_path is not None:
        request_validation = validate_delivery_request(request_path)
        if request_validation.get("status") != "PASS":
            failures.append("delivery request did not PASS validation for provenance expectations")
        else:
            request = request_validation["request"]
            request_plan = compile_delivery_plan(request_path)
            if request_plan.get("status") != "PASS":
                failures.append("delivery request did not compile to an accepted build definition for provenance expectations")
            else:
                expected_build_type = request_plan.get("buildType")
            expected_external["deliveryRequest"] = {
                "id": request.get("requestId"),
                "sha256": sha256_file(request_path),
            }
            expected_dependencies.append(_slsa_resource_descriptor(request_path))
            source = request_validation.get("resolved", {}).get("source")
            if isinstance(source, dict) and isinstance(source.get("path"), str):
                expected_dependencies.append(_slsa_resource_descriptor(Path(source["path"])))
            for material in request_validation.get("resolved", {}).get("materials", []):
                if isinstance(material, dict) and isinstance(material.get("path"), str):
                    expected_dependencies.append(_slsa_resource_descriptor(Path(material["path"])))
            if profile_path is not None:
                profile_ref = request.get("profile", {}) if isinstance(request.get("profile"), dict) else {}
                if profile_ref.get("id") != expected_external["deliveryProfile"]["id"]:
                    failures.append("delivery request profile id does not match provenance profile expectation")
                if profile_ref.get("sha256") != expected_external["deliveryProfile"]["sha256"]:
                    failures.append("delivery request profile digest does not match provenance profile expectation")
    if expected_build_type is not None and build_type != expected_build_type:
        failures.append("provenance buildType does not match the build definition derived from the delivery request")
    if expected_external and external_parameters != expected_external:
        failures.append("provenance externalParameters do not exactly match the expected delivery request/profile parameters")
    missing_dependencies = [item for item in expected_dependencies if item not in resolved_dependencies]
    if missing_dependencies:
        failures.append("provenance resolvedDependencies omit one or more digest-bound request/profile/material inputs")

    return {
        "status": "PASS" if not failures else "FAIL",
        "statement": file_fact(provenance_path),
        "expectedSubjects": expected,
        "observedSubjects": observed,
        "builderId": builder,
        "buildType": build_type,
        "expectedBuildType": expected_build_type,
        "externalParameters": external_parameters,
        "expectedExternalParameters": expected_external,
        "resolvedDependencies": resolved_dependencies,
        "expectedDependencies": expected_dependencies,
        "requestValidation": request_validation,
        "authenticity": "NOT_VERIFIED",
        "failures": failures,
        "boundary": "PASS validates SLSA Provenance v1 semantics and exact artifact/request/profile/buildType expectations only. builder.id is intentionally not tenant-requested; cryptographic authenticity requires a trusted signer-builder pair and signature/root-of-trust verification.",
    }


def verify_signed_release_provenance(
    statement_path: Path,
    bundle_path: Path,
    subjects: Iterable[Path],
    profile_path: Path,
    request_path: Path,
    trust_policy_path: Path,
    signer_id: str,
) -> dict[str, Any]:
    failures: list[str] = []
    subject_paths = list(subjects)
    if not subject_paths:
        failures.append("release provenance verification requires at least one subject artifact")
    policy_result = validate_attestation_trust_policy(trust_policy_path)
    if policy_result.get("status") != "PASS":
        failures.append("attestation trust policy did not PASS validation")
    signer = policy_result.get("resolvedSigners", {}).get(signer_id)
    if not isinstance(signer, dict):
        failures.append(f"trusted signer id is absent from policy: {signer_id}")
        signer = {}
    semantic = verify_release_provenance(
        statement_path,
        subject_paths,
        profile_path=profile_path,
        request_path=request_path,
    )
    if semantic.get("status") != "PASS":
        failures.append("SLSA Provenance semantic/request binding validation failed")
    builder_id = semantic.get("builderId")
    if signer and builder_id not in set(signer.get("allowedBuilderIds", [])):
        failures.append("SLSA Provenance builder.id is not authorized for the selected signer")
    accepted = policy_result.get("policy", {}).get("acceptedBundleMediaTypes", []) if isinstance(policy_result.get("policy"), dict) else []
    mode = str(signer.get("mode", ""))
    bundle_shape = verify_sigstore_attestation_bundle_shape(
        bundle_path, statement_path, accepted, mode
    ) if bundle_path.is_file() else {"status": "FAIL", "failures": ["Sigstore bundle is absent"]}
    if bundle_shape.get("status") != "PASS":
        failures.append("Sigstore bundle shape/statement binding validation failed")
    if signer.get("requireTransparencyLog") is True and bundle_shape.get("tlogEntryCount", 0) < 1:
        failures.append("trust policy requires transparency-log evidence but bundle has no tlog entry")

    tool = cosign_tool_fact()
    cosign_result: dict[str, Any] = {"status": "NOT_RUN", "failures": []}
    if not failures and subject_paths:
        cosign_result = _cosign_verify_standard_bundle(
            bundle_path,
            subject_paths[0],
            signer,
            tool,
            SLSA_PROVENANCE_V1,
        )
        failures.extend(cosign_result.get("failures", []))
    elif tool.get("status") != "PASS":
        failures.append("cosign verifier provenance/version/digest validation failed")
    return {
        "status": "PASS" if not failures else "FAIL",
        "authenticity": "VERIFIED" if not failures else "NOT_VERIFIED",
        "signerId": signer_id,
        "signerMode": mode,
        "builderId": builder_id,
        "trustPolicy": policy_result,
        "bundleShape": bundle_shape,
        "semantic": semantic,
        "cosign": cosign_result,
        "tool": tool,
        "failures": failures,
        "boundary": "PASS requires an exact SLSA Provenance v1 subject/request/profile/buildType binding, an authorized signer->builder.id pair, a standardized Sigstore v0.3 DSSE bundle, a provenance-verified digest/version-pinned Cosign verifier, and successful cosign verification with --check-claims=true. This authenticates the provenance statement but does not by itself assign a SLSA Build level or prove tenant isolation/control-plane generation guarantees.",
    }

def build_release_manifest(
    profile_path: Path,
    primary: Path,
    companions: Iterable[Path] = (),
    evidence: Iterable[Path] = (),
    provenance: Path | None = None,
) -> dict[str, Any]:
    profile_result = validate_profile(profile_path)
    profile = profile_result["profile"]
    failures: list[str] = []
    suffix_map = {"pptx": ".pptx", "docx": ".docx", "xlsx": ".xlsx", "html": ".html", "pdf": ".pdf", "pdf-a-4": ".pdf", "pdf-ua-2": ".pdf"}
    primary_format = str(profile.get("primaryOutput", {}).get("format", ""))
    expected_suffix = suffix_map.get(primary_format)
    if expected_suffix and primary.suffix.casefold() != expected_suffix:
        failures.append(f"primary artifact suffix {primary.suffix} does not match profile format {primary_format}")
    primary_fact = file_fact(primary)
    companion_paths = list(companions)
    declared_companions = list(profile.get("companions", []))
    required_formats = [str(item.get("format")) for item in declared_companions if item.get("required") is True]
    observed_suffixes = [path.suffix.casefold() for path in companion_paths]
    required_suffixes = [suffix_map.get(fmt) for fmt in required_formats]
    if len(companion_paths) < len(required_formats):
        failures.append(f"profile requires {len(required_formats)} companion artifact(s), received {len(companion_paths)}")
    for suffix in required_suffixes:
        if suffix is not None and suffix not in observed_suffixes:
            failures.append(f"required companion format with suffix {suffix} is absent")
    companion_facts = [file_fact(path) for path in companion_paths]
    evidence_facts = [file_fact(path) for path in evidence]
    provenance_result: dict[str, Any] | None = None
    if provenance is not None:
        provenance_result = verify_release_provenance(provenance, [primary, *companion_paths])
        if provenance_result.get("status") != "PASS":
            failures.extend(f"release provenance: {item}" for item in provenance_result.get("failures", []))
    elif profile.get("gates", {}).get("releaseProvenance") is True:
        failures.append("profile requires release provenance but no provenance statement was supplied")
    return {
        "schemaVersion": 1,
        "kind": "artifact-delivery-release-manifest",
        "status": "PASS" if profile_result.get("status") == "PASS" and not failures else "FAIL",
        "profile": {"id": profile.get("id"), "sha256": sha256_file(profile_path)},
        "primary": primary_fact,
        "companions": companion_facts,
        "evidence": evidence_facts,
        "provenance": provenance_result,
        "failures": failures,
        "boundary": "Manifest PASS establishes package assembly and exact digest/provenance binding only. Target rendering, visual acceptance, accessibility/conformance and destination read-back retain their independent gates.",
    }


def verify_render_evidence(render_dir: Path, expected_slides: int) -> dict[str, Any]:
    pngs = (
        sorted(p for p in render_dir.iterdir() if p.is_file() and p.suffix.casefold() == ".png")
        if render_dir.is_dir()
        else []
    )
    bad = [str(p) for p in pngs if p.stat().st_size <= 0]
    ok = len(pngs) == expected_slides and not bad and expected_slides > 0
    return {
        "status": "PASS" if ok else "FAIL",
        "renderDirectory": str(render_dir.resolve()),
        "expectedSlideCount": expected_slides,
        "observedPngCount": len(pngs),
        "emptyFiles": bad,
        "digests": [{"name": p.name, "sha256": sha256_file(p)} for p in pngs],
        "note": "Presence/count/digest evidence does not replace visual defect review.",
    }


def verify_target_evidence(
    evidence_path: Path,
    pptx: Path,
    pdf: Path | None,
    expected_slides: int,
    render_result: dict[str, Any] | None = None,
) -> dict[str, Any]:
    value = load_json(evidence_path)
    failures: list[str] = []
    expected_pptx = sha256_file(pptx)
    if value.get("artifact", {}).get("sha256") != expected_pptx:
        failures.append("target evidence PPTX digest does not match the inspected artifact")
    renderer = value.get("renderer", {})
    if renderer.get("name") != "Microsoft PowerPoint Desktop":
        failures.append("target evidence renderer is not Microsoft PowerPoint Desktop")
    if int(value.get("result", {}).get("slideCount", -1)) != expected_slides:
        failures.append("target evidence slide count mismatch")
    if pdf is not None:
        expected_pdf = sha256_file(pdf)
        if value.get("result", {}).get("pdfSha256") != expected_pdf:
            failures.append("target evidence PDF digest mismatch")
    if not value.get("renderer", {}).get("version"):
        failures.append("target evidence omitted PowerPoint version")
    if render_result is not None and render_result.get("status") == "PASS":
        expected_render_map = {item["name"]: item["sha256"] for item in render_result.get("digests", [])}
        target_pngs = value.get("result", {}).get("pngs")
        if not isinstance(target_pngs, list):
            failures.append("target evidence omitted PNG digest list")
        else:
            observed_render_map = {
                str(item.get("name")): str(item.get("sha256"))
                for item in target_pngs
                if isinstance(item, dict)
            }
            if observed_render_map != expected_render_map:
                failures.append("target evidence PNG digests do not match rendered evidence")
        if int(value.get("result", {}).get("pngCount", -1)) != expected_slides:
            failures.append("target evidence PNG count mismatch")
    return {
        "status": "PASS" if not failures else "FAIL",
        "evidencePath": str(evidence_path.resolve()),
        "renderer": renderer,
        "failures": failures,
    }


def verify_visual_review(evidence_path: Path, pptx: Path, render_result: dict[str, Any]) -> dict[str, Any]:
    value = load_json(evidence_path)
    failures: list[str] = []
    expected_artifact = sha256_file(pptx)
    if value.get("artifactSha256") != expected_artifact:
        failures.append("visual review artifact digest mismatch")
    if value.get("verdict") != "PASS":
        failures.append("visual review verdict is not PASS")
    blocking = value.get("blockingDefects")
    if not isinstance(blocking, list) or blocking:
        failures.append("visual review must provide an empty blockingDefects list")
    expected_renders = {item["name"]: item["sha256"] for item in render_result.get("digests", [])}
    observed_renders = value.get("renderDigests")
    if not isinstance(observed_renders, dict) or observed_renders != expected_renders:
        failures.append("visual review render digests do not bind the exact rendered evidence")
    methods = value.get("methods")
    if not isinstance(methods, list) or not methods or not all(isinstance(item, str) and item for item in methods):
        failures.append("visual review must name at least one review method")
    return {
        "status": "PASS" if not failures else "FAIL",
        "evidencePath": str(evidence_path.resolve()),
        "methods": methods if isinstance(methods, list) else [],
        "blockingDefects": blocking if isinstance(blocking, list) else None,
        "failures": failures,
        "boundary": "Visual-review evidence is bound to exact PPTX and render digests; structural/target/delivery validators remain independent gates.",
    }


def verify_readback(
    source: Path,
    readback: Path,
    destination: str | None = None,
    destination_reference: str | None = None,
    artifact_role: str | None = None,
) -> dict[str, Any]:
    source_fact = file_fact(source)
    read_fact = file_fact(readback)
    matched = source_fact["digest"]["sha256"] == read_fact["digest"]["sha256"]
    return {
        "status": "PASS" if matched else "FAIL",
        "destination": destination,
        "destinationReference": destination_reference,
        "artifactRole": artifact_role,
        "source": source_fact,
        "readback": read_fact,
        "digestMatched": matched,
    }


def verify_delivery_evidence(
    evidence_paths: Iterable[Path],
    profile: dict[str, Any],
    pptx: Path,
    pdf: Path | None,
) -> dict[str, Any]:
    expected_artifacts: dict[str, str] = {"primary": sha256_file(pptx)}
    required_companions = [
        item for item in profile.get("companions", []) if item.get("required") is True
    ]
    if required_companions:
        if pdf is None:
            return {
                "status": "FAIL",
                "failures": ["required companion delivery cannot be verified without a companion PDF"],
                "evidence": [],
            }
        expected_artifacts["companion"] = sha256_file(pdf)
    expected_pairs = {
        (str(destination), role)
        for destination in profile.get("deliveryTargets", [])
        for role in expected_artifacts
    }
    seen: set[tuple[str, str]] = set()
    failures: list[str] = []
    evidence: list[dict[str, Any]] = []
    for path in evidence_paths:
        value = load_json(path)
        destination = value.get("destination")
        role = value.get("artifactRole")
        pair = (str(destination), str(role))
        item_failures: list[str] = []
        if pair not in expected_pairs:
            item_failures.append(f"unexpected destination/artifactRole pair: {pair}")
        elif pair in seen:
            item_failures.append(f"duplicate destination/artifactRole evidence: {pair}")
        else:
            seen.add(pair)
        if value.get("status") != "PASS" or value.get("digestMatched") is not True:
            item_failures.append("read-back evidence did not PASS exact digest comparison")
        expected_digest = expected_artifacts.get(str(role))
        if expected_digest is not None and value.get("source", {}).get("digest", {}).get("sha256") != expected_digest:
            item_failures.append("read-back evidence source digest does not match the current build artifact")
        if not value.get("destinationReference"):
            item_failures.append("read-back evidence omitted destinationReference")
        failures.extend(f"{path}: {item}" for item in item_failures)
        evidence.append({"path": str(path.resolve()), "destination": destination, "artifactRole": role, "failures": item_failures})
    missing = sorted(expected_pairs - seen)
    if missing:
        failures.append("missing required destination/artifactRole read-back evidence: " + repr(missing))
    return {
        "status": "PASS" if not failures else "FAIL",
        "expected": sorted(expected_pairs),
        "evidence": evidence,
        "failures": failures,
        "boundary": "The delivery adapter owns the external write/read operation; this gate verifies exact returned bytes plus destination reference without inventing a transport protocol.",
    }


def snapshot_materials(paths: Iterable[Path]) -> dict[str, Any]:
    materials = [file_fact(path) for path in paths]
    return {
        "capturedAt": utc_now(),
        "materials": materials,
        "note": "Operational immutable-input snapshot. Durable build provenance should be emitted as an in-toto Statement/SLSA Provenance predicate.",
    }


def _require_uri(value: str, field: str) -> str:
    parsed = urllib.parse.urlparse(value)
    if not parsed.scheme:
        raise RuntimeError(f"{field} must be a URI")
    return value


def slsa_statement(
    subjects: Iterable[Path],
    materials: Iterable[Path],
    profile_path: Path,
    builder_id: str,
    build_type: str,
    request_path: Path | None = None,
) -> dict[str, Any]:
    builder_id = _require_uri(builder_id, "builder-id")
    build_type = _require_uri(build_type, "build-type")
    profile = load_json(profile_path)
    subject_values = [
        {"name": path.name, "digest": {"sha256": sha256_file(path)}} for path in subjects
    ]
    external_parameters: dict[str, Any] = {
        "deliveryProfile": {"id": profile.get("id"), "sha256": sha256_file(profile_path)},
    }
    dependency_paths = [*list(materials), profile_path]
    if request_path is not None:
        request = load_json(request_path)
        external_parameters["deliveryRequest"] = {
            "id": request.get("requestId"),
            "sha256": sha256_file(request_path),
        }
        dependency_paths.append(request_path)
    dependencies: list[dict[str, Any]] = []
    seen_dependencies: set[tuple[str, str, str]] = set()
    for path in dependency_paths:
        descriptor = _slsa_resource_descriptor(path)
        key = (descriptor["uri"], descriptor["name"], descriptor["digest"]["sha256"])
        if key not in seen_dependencies:
            seen_dependencies.add(key)
            dependencies.append(descriptor)
    return {
        "_type": IN_TOTO_STATEMENT_V1,
        "subject": subject_values,
        "predicateType": SLSA_PROVENANCE_V1,
        "predicate": {
            "buildDefinition": {
                "buildType": build_type,
                "externalParameters": external_parameters,
                "internalParameters": {},
                "resolvedDependencies": dependencies,
            },
            "runDetails": {
                "builder": {"id": builder_id},
                "metadata": {"invocationId": f"artifact-build-{utc_now()}"},
                "byproducts": [],
            },
        },
    }

def presentation_gate(
    profile_path: Path,
    pptx: Path,
    pdf: Path | None,
    render_dir: Path | None,
    openxml_evidence: Path | None,
    target_evidence: Path | None,
    visual_evidence: Path | None,
    delivery_evidence: Iterable[Path],
    font_dir: Path,
) -> dict[str, Any]:
    profile_result = validate_profile(profile_path)
    profile = profile_result["profile"]
    gates = profile.get("gates", {})
    placeholders = profile.get("semanticPolicy", {}).get("placeholderPatterns", [])
    pptx_result = inspect_pptx(pptx, placeholders)
    slide_count = int(pptx_result.get("package", {}).get("slideCount", 0))
    if openxml_evidence is None:
        openxml_result = {"status": "NOT_RUN", "reason": "no DocumentFormat.OpenXml validation evidence supplied"}
    else:
        openxml_result = verify_openxml_evidence(openxml_evidence, pptx)
    if pptx_result.get("status") != "PASS":
        structural_status = "FAIL"
    elif openxml_result.get("status") != "PASS":
        structural_status = openxml_result.get("status", "FAIL")
    else:
        structural_status = "PASS"
    semantic_result = verify_presentation_semantics(profile, pptx_result)
    font_result = verify_font_manifest(
        profile,
        font_dir,
        pptx_result.get("package", {}).get("renderExplicitTypefaceNames", []),
    )
    pdf_result: dict[str, Any]
    if pdf is None:
        pdf_result = {"status": "NOT_RUN", "reason": "no companion PDF supplied"}
    else:
        pdf_result = verify_pdf(pdf)
    render_result: dict[str, Any]
    if render_dir is None:
        render_result = {"status": "NOT_RUN", "reason": "no rendered PNG directory supplied"}
    else:
        render_result = verify_render_evidence(render_dir, slide_count)
    target_result: dict[str, Any]
    if target_evidence is None:
        target_result = {"status": "NOT_RUN", "reason": "no target PowerPoint evidence supplied"}
    else:
        target_result = verify_target_evidence(target_evidence, pptx, pdf, slide_count, render_result)
    visual_result: dict[str, Any]
    if visual_evidence is None:
        visual_result = {"status": "NOT_RUN", "reason": "no digest-bound visual review evidence supplied"}
    elif render_result.get("status") != "PASS":
        visual_result = {"status": "FAIL", "reason": "render evidence integrity failed before visual review"}
    else:
        visual_result = verify_visual_review(visual_evidence, pptx, render_result)
    delivery_paths = list(delivery_evidence)
    if not delivery_paths:
        delivery_result = {"status": "NOT_RUN", "reason": "no destination read-back evidence supplied"}
    else:
        delivery_result = verify_delivery_evidence(delivery_paths, profile, pptx, pdf)

    component = {
        "profileSchema": profile_result["status"],
        "structural": structural_status,
        "dependency": font_result["status"],
        "semantic": semantic_result["status"],
        "companionPdf": pdf_result["status"],
        "visual": visual_result["status"],
        "target": target_result["status"],
        "deliveryReadback": delivery_result["status"],
    }
    required_failures = [
        name
        for name, required in gates.items()
        if required and (name not in component or component[name] != "PASS")
    ]
    return {
        "status": "PASS" if not required_failures else "FAIL",
        "profileId": profile.get("id"),
        "artifact": file_fact(pptx),
        "requiredGateFailures": required_failures,
        "components": {
            "profile": profile_result,
            "pptx": pptx_result,
            "openXmlValidation": openxml_result,
            "semantic": semantic_result,
            "fonts": font_result,
            "pdf": pdf_result,
            "renders": render_result,
            "visualReview": visual_result,
            "target": target_result,
            "deliveryReadback": delivery_result,
        },
        "truthBoundary": "PASS requires every profile-required gate represented here, including target PowerPoint evidence, visual review, and destination read-back. External delivery adapters still own the actual write/read effects.",
    }


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")


def emit(value: Any, output: Path | None) -> None:
    if output is not None:
        write_json(output, value)
    print(json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False))


def _parse_named_values(items: Iterable[str], option_name: str, path_values: bool = False) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for item in items:
        if "=" not in item:
            raise RuntimeError(f"{option_name} requires name=value")
        name, value = item.split("=", 1)
        if not name or not value or name in result:
            raise RuntimeError(f"duplicate/invalid {option_name} name: {name!r}")
        result[name] = Path(value) if path_values else value
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("validate-profile")
    p.add_argument("profile", type=Path)
    p.add_argument("--schema", type=Path, default=DEFAULT_SCHEMA)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("validate-request")
    p.add_argument("request", type=Path)
    p.add_argument("--schema", type=Path, default=DEFAULT_REQUEST_SCHEMA)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("compile-request")
    p.add_argument("request", type=Path)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("build-request")
    p.add_argument("request", type=Path)
    p.add_argument("--output-directory", type=Path)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("validate-presentation-source")
    p.add_argument("source", type=Path)
    p.add_argument("--schema", type=Path, default=DEFAULT_PRESENTATION_SOURCE_SCHEMA)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("build-presentation-source")
    p.add_argument("--source", type=Path, required=True)
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--pptx", type=Path, required=True)
    p.add_argument("--source-schema", type=Path, default=DEFAULT_PRESENTATION_SOURCE_SCHEMA)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("aggregate-gates")
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--gate", action="append", default=[], help="gateName=path/to/evidence.json")
    p.add_argument("--output", type=Path)

    p = sub.add_parser("verify-stage")
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--artifact", type=Path, required=True)
    p.add_argument("--output-directory", type=Path, required=True)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("verify-vsa")
    p.add_argument("--statement", type=Path, required=True)
    p.add_argument("--subject", type=Path, required=True)
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--allowed-verifier", action="append", default=[])
    p.add_argument("--output", type=Path)

    p = sub.add_parser("validate-attestation-trust-policy")
    p.add_argument("policy", type=Path)
    p.add_argument("--schema", type=Path, default=DEFAULT_ATTESTATION_TRUST_POLICY_SCHEMA)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("keyless-signing-readiness")
    p.add_argument("--trust-policy", type=Path, required=True)
    p.add_argument("--signer-id", required=True)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("verify-signed-vsa")
    p.add_argument("--statement", type=Path, required=True)
    p.add_argument("--bundle", type=Path, required=True)
    p.add_argument("--subject", type=Path, required=True)
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--trust-policy", type=Path, required=True)
    p.add_argument("--signer-id", required=True)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("verify-provenance")
    p.add_argument("--statement", type=Path, required=True)
    p.add_argument("--subject", type=Path, action="append", required=True)
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--request", type=Path)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("verify-signed-provenance")
    p.add_argument("--statement", type=Path, required=True)
    p.add_argument("--bundle", type=Path, required=True)
    p.add_argument("--subject", type=Path, action="append", required=True)
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--request", type=Path, required=True)
    p.add_argument("--trust-policy", type=Path, required=True)
    p.add_argument("--signer-id", required=True)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("aggregate-vsa-gates")
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--subject", type=Path, required=True)
    p.add_argument("--gate", action="append", default=[], help="gateName=path/to/vsa.json")
    p.add_argument("--bundle", action="append", default=[], help="gateName=path/to/sigstore-bundle.json")
    p.add_argument("--signer", action="append", default=[], help="gateName=trusted-signer-id")
    p.add_argument("--trust-policy", type=Path)
    p.add_argument("--allow-local-unsigned", action="store_true")
    p.add_argument("--output", type=Path)

    p = sub.add_parser("package-stage")
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--primary", type=Path, required=True)
    p.add_argument("--verify-report", type=Path, required=True)
    p.add_argument("--output-directory", type=Path, required=True)
    p.add_argument("--companion", type=Path, action="append", default=[])
    p.add_argument("--request", type=Path)
    p.add_argument("--provenance", type=Path)
    p.add_argument("--provenance-bundle", type=Path)
    p.add_argument("--provenance-signer")
    p.add_argument("--gate-bundle", action="append", default=[], help="gateName=path/to/sigstore-bundle.json")
    p.add_argument("--gate-signer", action="append", default=[], help="gateName=trusted-signer-id")
    p.add_argument("--trust-policy", type=Path)
    p.add_argument("--allow-local-unsigned", action="store_true")
    p.add_argument("--output", type=Path)

    p = sub.add_parser("inspect-pptx")
    p.add_argument("pptx", type=Path)
    p.add_argument("--placeholder", action="append", default=[])
    p.add_argument("--output", type=Path)

    p = sub.add_parser("verify-pdf")
    p.add_argument("pdf", type=Path)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("verify-pdf-conformance")
    p.add_argument("pdf", type=Path)
    p.add_argument("--flavour", required=True, choices=("4", "4f", "4e", "ua1", "ua2", "wt1r", "wt1a"))
    p.add_argument("--output", type=Path)

    p = sub.add_parser("release-manifest")
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--primary", type=Path, required=True)
    p.add_argument("--companion", type=Path, action="append", default=[])
    p.add_argument("--evidence", type=Path, action="append", default=[])
    p.add_argument("--provenance", type=Path)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("snapshot")
    p.add_argument("material", type=Path, nargs="+")
    p.add_argument("--output", type=Path, required=True)

    p = sub.add_parser("verify-readback")
    p.add_argument("source", type=Path)
    p.add_argument("readback", type=Path)
    p.add_argument("--destination", required=True)
    p.add_argument("--destination-reference", required=True)
    p.add_argument("--artifact-role", choices=("primary", "companion"), required=True)
    p.add_argument("--output", type=Path)

    p = sub.add_parser("attest-slsa")
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--subject", type=Path, action="append", required=True)
    p.add_argument("--material", type=Path, action="append", default=[])
    p.add_argument("--builder-id", required=True)
    p.add_argument("--build-type", required=True)
    p.add_argument("--request", type=Path)
    p.add_argument("--output", type=Path, required=True)

    p = sub.add_parser("gate-presentation")
    p.add_argument("--profile", type=Path, required=True)
    p.add_argument("--pptx", type=Path, required=True)
    p.add_argument("--pdf", type=Path)
    p.add_argument("--render-dir", type=Path)
    p.add_argument("--openxml-evidence", type=Path)
    p.add_argument("--target-evidence", type=Path)
    p.add_argument("--visual-evidence", type=Path)
    p.add_argument("--delivery-evidence", type=Path, action="append", default=[])
    p.add_argument("--font-dir", type=Path, default=DEFAULT_WINDOWS_FONTS)
    p.add_argument("--output", type=Path)

    args = parser.parse_args()
    try:
        if args.command == "validate-profile":
            value = validate_profile(args.profile, args.schema)
        elif args.command == "validate-request":
            value = validate_delivery_request(args.request, args.schema)
        elif args.command == "compile-request":
            value = compile_delivery_plan(args.request)
        elif args.command == "build-request":
            value = execute_build_stage(args.request, args.output_directory)
        elif args.command == "validate-presentation-source":
            value = validate_json_document(args.source, args.schema, "presentation-source")
        elif args.command == "build-presentation-source":
            value = build_presentation_source(args.source, args.profile, args.pptx, args.source_schema)
        elif args.command == "aggregate-gates":
            gate_paths: dict[str, Path] = {}
            for item in args.gate:
                if "=" not in item:
                    raise RuntimeError("--gate requires gateName=path")
                gate, path_text = item.split("=", 1)
                if not gate or gate in gate_paths:
                    raise RuntimeError(f"duplicate/invalid gate name: {gate!r}")
                gate_paths[gate] = Path(path_text)
            value = aggregate_gate_results(args.profile, gate_paths)
        elif args.command == "verify-stage":
            value = execute_verify_stage(args.profile, args.artifact, args.output_directory)
        elif args.command == "verify-vsa":
            value = verify_verification_summary(args.statement, args.subject, args.profile, args.allowed_verifier)
        elif args.command == "validate-attestation-trust-policy":
            value = validate_attestation_trust_policy(args.policy, args.schema)
        elif args.command == "keyless-signing-readiness":
            value = production_keyless_signing_readiness(args.trust_policy, args.signer_id)
        elif args.command == "verify-signed-vsa":
            value = verify_signed_verification_summary(
                args.statement, args.bundle, args.subject, args.profile, args.trust_policy, args.signer_id
            )
        elif args.command == "verify-provenance":
            value = verify_release_provenance(
                args.statement, args.subject, profile_path=args.profile, request_path=args.request
            )
        elif args.command == "verify-signed-provenance":
            value = verify_signed_release_provenance(
                args.statement,
                args.bundle,
                args.subject,
                args.profile,
                args.request,
                args.trust_policy,
                args.signer_id,
            )
        elif args.command == "aggregate-vsa-gates":
            gate_paths = _parse_named_values(args.gate, "--gate", path_values=True)
            bundle_paths = _parse_named_values(args.bundle, "--bundle", path_values=True)
            signer_ids = _parse_named_values(args.signer, "--signer")
            value = aggregate_vsa_gates(
                args.profile,
                args.subject,
                gate_paths,
                args.allow_local_unsigned,
                bundle_paths,
                args.trust_policy,
                signer_ids,
            )
        elif args.command == "package-stage":
            gate_bundles = _parse_named_values(args.gate_bundle, "--gate-bundle", path_values=True)
            gate_signers = _parse_named_values(args.gate_signer, "--gate-signer")
            value = execute_package_stage(
                args.profile,
                args.primary,
                args.verify_report,
                args.output_directory,
                args.companion,
                args.request,
                args.provenance,
                args.provenance_bundle,
                args.provenance_signer,
                args.allow_local_unsigned,
                gate_bundles,
                args.trust_policy,
                gate_signers,
            )
        elif args.command == "inspect-pptx":
            value = inspect_pptx(args.pptx, args.placeholder)
        elif args.command == "verify-pdf":
            value = verify_pdf(args.pdf)
        elif args.command == "verify-pdf-conformance":
            value = verify_pdf_conformance(args.pdf, args.flavour)
        elif args.command == "release-manifest":
            value = build_release_manifest(args.profile, args.primary, args.companion, args.evidence, args.provenance)
        elif args.command == "snapshot":
            value = snapshot_materials(args.material)
        elif args.command == "verify-readback":
            value = verify_readback(
                args.source,
                args.readback,
                args.destination,
                args.destination_reference,
                args.artifact_role,
            )
        elif args.command == "attest-slsa":
            value = slsa_statement(
                args.subject, args.material, args.profile, args.builder_id, args.build_type, args.request
            )
        elif args.command == "gate-presentation":
            value = presentation_gate(
                args.profile,
                args.pptx,
                args.pdf,
                args.render_dir,
                args.openxml_evidence,
                args.target_evidence,
                args.visual_evidence,
                args.delivery_evidence,
                args.font_dir,
            )
        else:  # pragma: no cover
            raise AssertionError(args.command)
        emit(value, getattr(args, "output", None))
        status = value.get("status") if isinstance(value, dict) else None
        if args.command == "attest-slsa" or args.command == "snapshot":
            return 0
        return 0 if status == "PASS" else 1
    except Exception as error:
        print(json.dumps({"status": "ERROR", "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
